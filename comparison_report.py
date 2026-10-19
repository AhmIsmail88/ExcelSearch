"""Professional, streaming Excel reports for immutable comparison snapshots."""
import json
import math
import sqlite3
from datetime import datetime, timezone, timedelta

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


LABELS={'suggested':'اقتراح للمراجعة','accepted':'مقبول يدويًا','rejected':'مرفوض',
        'changed':'مختلف','identical':'متطابق','only_a':'فقط في الأول ضمن النطاق','only_b':'فقط في الثاني ضمن النطاق'}
COLORS={'suggested':'FFF2CC','accepted':'DDEFE5','rejected':'FCE4D6',
        'changed':'FFF2CC','identical':'DDEFE5','only_a':'E3EDF7','only_b':'EEE7F5'}
INK='19354A';BRASS='A97C2F';MUTED='607587';PALE='F2F5F7'


def export_comparison(engine, destination):
    wb=Workbook(write_only=True)
    wb.properties.title='تقرير مقارنة البنود'
    wb.properties.subject='نتائج المقارنة ونطاق الفلترة وقرارات المراجعة'
    wb.properties.creator='أرشيف البحث'
    wb.properties.language='ar-SA'
    summary=wb.create_sheet('ملخص المقارنة')
    details=wb.create_sheet('المقارنة')
    widths=[24,16,12,12,24,30,65,65]
    row_counts={}

    def setup(sheet, columns, title, freeze='A6'):
        sheet.sheet_view.rightToLeft=True
        sheet.sheet_view.showGridLines=False
        sheet.sheet_properties.pageSetUpPr.fitToPage=True
        sheet.sheet_properties.outlinePr.summaryRight=False
        sheet.sheet_properties.tabColor=BRASS
        sheet.page_setup.orientation='landscape'
        sheet.page_setup.paperSize=Worksheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
        sheet.print_options.horizontalCentered=True
        sheet.oddFooter.center.text='صفحة &P من &N'
        sheet.oddFooter.right.text='أرشيف البحث | مقارنة البنود'
        for col,width in enumerate(columns,1):
            sheet.column_dimensions[get_column_letter(col)].width=width
        sheet.freeze_panes=freeze
        sheet.print_title_rows='1:5'
        sheet.merged_cells.add(f'A1:{get_column_letter(len(columns))}1')
        sheet.row_dimensions[1].height=36
        sheet.row_dimensions[2].height=28
        sheet.row_dimensions[3].height=28
        append(sheet,[title]+['']*(len(columns)-1),fill=INK,color='FFFFFF',bold=True,size=19)

    def append(sheet,values,fill=None,color=INK,bold=False,size=11,status=None,changed=False,formats=None,highlights=None):
        row_number=row_counts.get(sheet.title,0)+1
        row_counts[sheet.title]=row_number
        if sheet is summary:
            lengths=[28,62,28,62]
            lines=max(1,*[sum(max(1,math.ceil(len(part)/lengths[i])) for part in str(value or '').split('\n')) for i,value in enumerate(values)])
            sheet.row_dimensions[row_number].height=max(sheet.row_dimensions[row_number].height or 0,24,lines*16)
        cells=[]
        for index,value in enumerate(values):
            cell=WriteOnlyCell(sheet,value=value)
            if isinstance(value,str):
                cell.data_type='s'
            cell.font=Font(name='Arial',size=size,color=color,bold=bold)
            cell.alignment=Alignment(horizontal='right',vertical='top',wrap_text=True,readingOrder=2)
            cell.border=Border(bottom=Side(style='hair',color='DDE4E8'))
            cell.fill=PatternFill('solid',fgColor=fill or 'FFFFFF')
            if status and index==0:
                cell.fill=PatternFill('solid',fgColor=COLORS[status]);cell.font=Font(name='Arial',size=11,bold=True,color=INK)
            if status and index==1 and isinstance(value,(int,float)):
                cell.number_format='0.0%'
                cell.alignment=Alignment(horizontal='center',vertical='top')
            if changed and index in (6,7):
                cell.fill=PatternFill('solid',fgColor='FFF1ED' if index==6 else 'EDF6F0')
            if formats and index in formats and isinstance(value,(int,float)):
                cell.number_format=formats[index]
                cell.alignment=Alignment(horizontal='center',vertical='top')
            if highlights and index in highlights:
                cell.fill=PatternFill('solid',fgColor='FFF2CC')
            cells.append(cell)
        sheet.append(cells)
        if sheet is not summary and row_number in sheet.row_dimensions:
            del sheet.row_dimensions[row_number]

    conn=sqlite3.connect(engine.path)
    try:
        counts=engine.counts(conn)
        setup(summary,[28,62,28,62],'تقرير مقارنة البنود | ملخص المراجعة')
        generated=datetime.now(timezone(timedelta(hours=3))).strftime('%Y-%m-%d %H:%M')
        append(summary,['وقت التصدير (الرياض)',generated,'طريقة المقارنة','تشابه الوصف مع مراجعة يدوية' if engine.mode=='similarity' else 'تطابق المفتاح حرفيًا'],fill=PALE)
        append(summary,['الملف الأول',engine.names['a'],'الملف الثاني',engine.names['b']])
        append(summary,['عمود الربط / الوصف',engine.summary['key_a'],'العمود المقابل',engine.summary['key_b']],fill=PALE)
        append(summary,['المؤشر','الملف الأول','المؤشر','الملف الثاني'],fill=INK,color='FFFFFF',bold=True)
        append(summary,['صفوف المصدر',engine.summary['source_count_a'],'صفوف المصدر',engine.summary['source_count_b']])
        append(summary,['بعد الفلترة',engine.summary['filtered_count_a'],'بعد الفلترة',engine.summary['filtered_count_b']],fill=PALE)
        append(summary,['مفتاح / وصف فارغ',engine.summary['blank_a'],'مفتاح / وصف فارغ',engine.summary['blank_b']])
        append(summary,['كلمة البحث',engine.query or 'بدون','الحد الأدنى للتشابه',str(engine.summary.get('threshold','—'))],fill=PALE)
        append(summary,['عمود القيمة المختار','أعمدة الكميات المحددة' if engine.quantity_mode else engine.summary.get('field_a') or 'كل الأعمدة المشتركة','العمود المقابل','حسب ربط السنوات المحدد' if engine.quantity_mode else engine.summary.get('field_b') or 'كل الأعمدة المشتركة'])
        append(summary,['حالة النتيجة','عدد العناصر','حالة النتيجة','عدد العناصر'],fill=INK,color='FFFFFF',bold=True)
        relevant=['suggested','accepted','rejected'] if engine.mode=='similarity' else ['changed','identical','only_a','only_b']
        for i in range(0,len(relevant),2):
            left=relevant[i];right=relevant[i+1] if i+1<len(relevant) else None
            append(summary,[LABELS[left],counts.get(left,0),LABELS[right] if right else '',counts.get(right,0) if right else ''],fill=PALE)
        if getattr(engine,'has_qty',False):
            qcounts=engine.qty_counts(conn)
            append(summary,['بنود بها فرق كميات',sum(qcounts.get(section,0) for section in relevant),'','الفرق = قيمة الثاني ناقص الأول؛ تُقارن الكميات سنة بسنة مع الإجمالي.'],fill=PALE)
        append(summary,['نطاق الفلترة','الشرط','القيمة','ملاحظة'],fill=INK,color='FFFFFF',bold=True)
        operators={'contains':'يحتوي','not_contains':'لا يحتوي — استبعاد','equals':'يساوي','empty':'فارغ','not_empty':'غير فارغ'}
        for tag in ('a','b'):
            filters=engine.filters[tag]
            if not filters:
                append(summary,[engine.names[tag],'بدون فلاتر','',''])
            for item in filters:
                column='أي خلية في الصف' if item.get('scope')=='all' else item['col']
                append(summary,[engine.names[tag],column+' — '+operators[item.get('op','contains')],item.get('value',''),'كل الشروط مطلوبة'])
        append(summary,['طريقة القراءة','المقارنة مبنية على لقطة ثابتة وقت التنفيذ. النتائج تخص الصفوف التي اجتازت الفلاتر فقط.','',''],fill=PALE)
        append(summary,['النصوص الطويلة','النص محفوظ كاملًا. إذا تجاوز ارتفاع العرض، استخدم شريط الصيغة لقراءة الخلية كاملة. أرقام الصفوف هي معرفات داخل القاعدة وليست أرقام Excel الأصلية.','',''])
        if engine.mode=='similarity':
            append(summary,['حدود الاقتراحات','حتى 3 اقتراحات لكل صف من أفضل 25 مرشحًا نصيًا. النسبة ليست ضمانًا لتكافؤ البنود.','',''])
            append(summary,['بدون اقتراح في الأول',engine.summary.get('without_suggestion_a',0),'قرارات المراجعة','قبول أو رفض بشري؛ لا يعدّل المصدر'],fill=PALE)
        if engine.summary.get('snapshot_created_at'):
            append(summary,['وقت لقطة البيانات (UTC)',engine.summary['snapshot_created_at'],'نوع البيانات','لقطة محفوظة وقت تنفيذ المقارنة'])
        if len(engine.key_columns['a'])>1:
            append(summary,['أعمدة الربط في الأول',' + '.join(engine.key_columns['a']),'أعمدة الربط في الثاني',' + '.join(engine.key_columns['b'])])
        if conn.execute("SELECT 1 FROM results WHERE note<>'' LIMIT 1").fetchone():
            notes=wb.create_sheet('ملاحظات المراجعة');part=1;line=5
            def start_notes(sheet):
                setup(sheet,[24,16,16,16,30,80],'ملاحظات المراجعة المحفوظة')
                append(sheet,['الملف الأول',engine.names['a'],'الملف الثاني',engine.names['b']])
                append(sheet,['وقت التصدير',generated]);append(sheet,['الملاحظات مرتبطة بلقطة المقارنة'])
                append(sheet,['الحالة','التشابه','صف أ','صف ب','المفتاح / الزوج','الملاحظة'],fill=INK,color='FFFFFF',bold=True)
            start_notes(notes)
            for section,score,aid,bid,raw,note in conn.execute("SELECT section,score,aid,bid,body,note FROM results WHERE note<>'' ORDER BY id"):
                if line>=1048576:
                    notes.auto_filter.ref=f'A5:F{line}';part+=1;notes=wb.create_sheet(f'ملاحظات {part}');start_notes(notes);line=5
                line+=1;notes.row_dimensions[line].height=min(390,max(40,16*math.ceil(len(note)/70)))
                append(notes,[LABELS[section],None if score is None else score/100,aid,bid,json.loads(raw)['key'],note],status=section)
            notes.auto_filter.ref=f'A5:F{line}'
        if engine.quantity_mode:
            from quantity_comparison import MODES
            from quantity_report import write_quantity_report
            append(summary,['مقارنة الكميات',MODES[engine.quantity_mode],'قاعدة الإجمالي','جمع القيم الرقمية الموجودة؛ الفراغ لا يساوي الصفر'],fill=PALE)
            append(summary,['عمود كمية الأول','عمود كمية الثاني','طريقة الحساب',''],fill=INK,color='FFFFFF',bold=True)
            for ca,cb in engine.qty_pairs:
                append(summary,[ca,cb,'الثاني ناقص الأول',''])
            append(summary,['أعمدة العرض في الأول','، '.join(engine.display_columns['a']),'أعمدة العرض في الثاني','، '.join(engine.display_columns['b'])])
            write_quantity_report(engine,conn,wb,details,setup,append)
            wb.save(destination)
            return
        setup(details,widths,'مقارنة البنود | تفاصيل النتائج')
        append(details,['الملف الأول',engine.names['a'],'','الملف الثاني',engine.names['b'],'','',''],fill=PALE)
        append(details,['وقت التصدير',generated,'','نطاق البحث',engine.query or 'كل الصفوف بعد الفلاتر','','',''])
        append(details,['دليل الألوان','أصفر: مراجعة / اختلاف','أخضر: قبول','برتقالي: رفض','','القيم المختلفة مميزة لونيًا','',''],fill=PALE)
        headers=['الحالة','التشابه النصي','صف أ بالقاعدة','صف ب بالقاعدة','المفتاح / الزوج','الحقل',engine.names['a'],engine.names['b']]
        details.row_dimensions[5].height=32
        append(details,headers,fill=INK,color='FFFFFF',bold=True)
        line=5;part=1
        for section,aid,bid,score,raw in conn.execute('SELECT section,aid,bid,score,body FROM results ORDER BY id'):
            body=json.loads(raw)
            pairs=list(dict.fromkeys(engine.field_pairs+[(c,c) for c in dict.fromkeys([*body['row_a'],*body['row_b']]) if c not in {x for pair in engine.field_pairs for x in pair}]))
            if engine.mode=='key' and engine.summary.get('field_a'):
                pairs=list(dict.fromkeys([(engine.summary['key_a'],engine.summary['key_b'])]+engine.field_pairs))
            for ca,cb in pairs:
                if line>=1048576:
                    details.auto_filter.ref=f'A5:H{line}'
                    part+=1;details=wb.create_sheet(f'المقارنة {part}')
                    setup(details,widths,'مقارنة البنود | تابع النتائج')
                    append(details,['الملف الأول',engine.names['a'],'','الملف الثاني',engine.names['b'],'','',''])
                    append(details,['وقت التصدير',generated,'','','','','','']);append(details,['']*8)
                    append(details,headers,fill=INK,color='FFFFFF',bold=True);line=5
                va,vb=body['row_a'].get(ca,''),body['row_b'].get(cb,'')
                line+=1
                lines=max(1,*[sum(max(1,math.ceil(len(chunk)/55)) for chunk in str(value or '').split('\n')) for value in (va,vb)])
                details.row_dimensions[line].height=min(390,max(34,lines*15))
                append(details,[LABELS[section],None if score is None else score/100,aid,bid,body['key'],ca if ca==cb else ca+' ↔ '+cb,va,vb],
                       fill='F6F8FA' if line%2==0 else 'FFFFFF',status=section,changed=(ca if ca==cb else ca+' ↔ '+cb) in body['diffs'])
        details.auto_filter.ref=f'A5:H{max(line,5)}'
        if getattr(engine,'has_qty',False):
            qcols=[(ca or cb) for ca,cb in engine.qty_pairs]
            qn=len(qcols)
            qsums={'a':[0.0]*qn,'b':[0.0]*qn};qfilled={'a':[0]*qn,'b':[0]*qn}
            for section,aid,bid,score,raw in conn.execute('SELECT section,aid,bid,score,body FROM results ORDER BY id'):
                body=json.loads(raw);qty=body.get('qty')
                if not qty: continue
                # Suggestions can share source rows. Only reviewed one-to-one
                # matches contribute to similarity totals.
                if engine.mode=='similarity' and section!='accepted': continue
                for index,c in enumerate(qty['cols'][:qn]):
                    if c.get('a_num') is not None:
                        qsums['a'][index]+=c['a_num'];qfilled['a'][index]+=1
                    if c.get('b_num') is not None:
                        qsums['b'][index]+=c['b_num'];qfilled['b'][index]+=1
            qsheet=wb.create_sheet('الكميات')
            setup(qsheet,[30,18,16,16,18,18,40],'مقارنة البنود | الكميات')
            append(qsheet,['الإجماليات للتطابقات المقبولة يدويًا فقط؛ تفاصيل الاقتراحات أدناه لا تدخل الإجمالي.' if engine.mode=='similarity' else 'الفرق = قيمة الثاني ناقص الأول. الفراغ يعني «لا كمية» ولا يُحسب ضمن الإجمالي.','','','','','',''],fill=PALE)
            append(qsheet,['العمود','إجمالي الأول','إجمالي الثاني','الفرق','عدد القيم في الأول','عدد القيم في الثاني',''],fill=INK,color='FFFFFF',bold=True)
            for index,name in enumerate(qcols):
                append(qsheet,[name,round(qsums['a'][index],6),round(qsums['b'][index],6),
                               round(qsums['b'][index]-qsums['a'][index],6),qfilled['a'][index],qfilled['b'][index],''],
                       fill=PALE if index%2==0 else None,formats={1:'#,##0.###',2:'#,##0.###',3:'#,##0.###',4:'0',5:'0'})
            append(qsheet,['']*7)
            append(qsheet,['الحالة','الزوج / المفتاح','العمود','قيمة الأول','قيمة الثاني','الفرق','ملاحظة'],fill=INK,color='FFFFFF',bold=True)
            for section,raw in conn.execute('SELECT section,body FROM results ORDER BY id'):
                body=json.loads(raw);qty=body.get('qty');key=body.get('key','')
                if not qty: continue
                for c in qty['cols'][:qn]:
                    va=c['a_num'] if c.get('a_num') is not None else (c.get('a') or '')
                    vb=c['b_num'] if c.get('b_num') is not None else (c.get('b') or '')
                    append(qsheet,[LABELS.get(section,section),key,c.get('name',''),va,vb,
                                   c.get('delta') if c.get('delta') is not None else '',''],
                           formats={3:'#,##0.###',4:'#,##0.###',5:'#,##0.###'})
                append(qsheet,['','','الإجمالي',qty['a_total'],qty['b_total'],qty['delta_total'],''],fill=PALE,
                       formats={3:'#,##0.###',4:'#,##0.###',5:'#,##0.###'})
        wb.save(destination)
    finally:
        conn.close()
