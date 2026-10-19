"""
excel_db_core.py
==================
المنطق الأساسي لتحويل ملف إكسل إلى قاعدة SQLite قابلة للبحث السريع،
مع دعم المزامنة التزايدية (Incremental Sync) بدل إعادة البناء الكامل كل مرة.

هذا الملف مكتبة (لا يعمل مباشرة) — يُستخدم من excel_search_db.py (CLI) و app.py (واجهة الويب).
"""

import hashlib
import json
import os
import re
import sqlite3
import time
import tempfile
import unicodedata
from pathlib import Path

import pandas as pd

# ----------------------------------------------------------------------
# تطبيع النص العربي للبحث الذكي
# ----------------------------------------------------------------------
_ARABIC_DIACRITICS = re.compile(r"[\u0610-\u061A\u064B-\u065F\u06D6-\u06DC\u06DF-\u06E8\u06EA-\u06ED\u0670]")
_TATWEEL = re.compile(r"\u0640")


def normalize_arabic(text) -> str:
    if text is None:
        return ""
    text = str(text)
    text = unicodedata.normalize("NFKC", text)
    text = _ARABIC_DIACRITICS.sub("", text)
    text = _TATWEEL.sub("", text)
    text = text.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا")
    text = text.replace("ى", "ي")
    text = text.replace("ة", "ه")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def sanitize_col_name(name) -> str:
    name = str(name).strip()
    name = re.sub(r"\s+", "_", name)
    name = re.sub(r"[^\w\u0600-\u06FF]", "_", name)
    if not name:
        name = "col"
    if name[0].isdigit():
        name = f"c_{name}"
    return name


def _row_hash(values) -> str:
    joined = "\x1f".join("" if v is None else str(v) for v in values)
    return hashlib.sha1(joined.encode("utf-8")).hexdigest()


def _file_hash(path: Path) -> str:
    h = hashlib.sha1()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


# ----------------------------------------------------------------------
# إدارة إعدادات القاعدة (جدول _config)
# ----------------------------------------------------------------------
def _ensure_config_table(conn):
    conn.execute("CREATE TABLE IF NOT EXISTS _config (key TEXT PRIMARY KEY, value TEXT)")


def set_config(conn, key, value):
    _ensure_config_table(conn)
    conn.execute(
        "INSERT INTO _config (key, value) VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
        (key, json.dumps(value, ensure_ascii=False)),
    )


def get_config(conn, key, default=None):
    _ensure_config_table(conn)
    row = conn.execute("SELECT value FROM _config WHERE key=?", (key,)).fetchone()
    if row is None:
        return default
    try:
        return json.loads(row[0])
    except (TypeError, json.JSONDecodeError):
        return row[0]


def get_all_config(db_path):
    conn = sqlite3.connect(db_path)
    _ensure_config_table(conn)
    rows = conn.execute("SELECT key, value FROM _config").fetchall()
    conn.close()
    out = {}
    for k, v in rows:
        try:
            out[k] = json.loads(v)
        except (TypeError, json.JSONDecodeError):
            out[k] = v
    return out


# ----------------------------------------------------------------------
# كشف صف العناوين الحقيقي تلقائياً (لملفات فيها ترويسة/بيانات مشروع فوق الجدول)
# ----------------------------------------------------------------------
def detect_header_row(excel_path, sheet=None, max_scan=20):
    """يفحص أول max_scan صف في الشيت، ويحاول يلاقي صف العناوين الحقيقي.
    المنطق: صف العناوين الحقيقي عادة بيكون كل خلاياه متعبّية (مش فاضية زي صفوف
    العنوان/الميتاداتا اللي بتبقى متفرقة)، وكل القيم فيه فريدة (مفيش تكرار
    زي البيانات العادية)."""
    try:
        raw = pd.read_excel(excel_path, sheet_name=sheet or 0, header=None, nrows=max_scan, dtype=str)
    except Exception:
        return 0

    n_cols = raw.shape[1]
    if n_cols == 0:
        return 0

    best_row, best_score = 0, -1.0
    for i in range(len(raw)):
        row = raw.iloc[i]
        mask = row.notna() & (row.astype(str).str.strip() != "")
        non_null_count = int(mask.sum())
        if non_null_count == 0:
            continue
        values = row[mask].astype(str).str.strip()
        unique_ratio = values.nunique() / non_null_count
        fill_ratio = non_null_count / n_cols
        score = fill_ratio * unique_ratio
        if score > best_score and fill_ratio >= 0.5 and unique_ratio >= 0.9:
            best_score = score
            best_row = i
    return best_row


# ----------------------------------------------------------------------
# قراءة وتنظيف ملف الإكسل (خطوة مشتركة بين build و sync)
# ----------------------------------------------------------------------
def _worksheet_records(ws, header_idx, compact_columns=True, last_column=None):
    """Project occupied columns, preserving values even far beyond the table."""
    ws.reset_dimensions()
    occupied = set()
    for row in ws.iter_rows(min_row=header_idx + 1, max_col=last_column, values_only=True):
        occupied.update(index for index, value in enumerate(row)
                        if value is not None and str(value).strip())
    width = max(occupied, default=-1) + 1
    indices = sorted(occupied) if compact_columns else list(range(width))
    records = ws.iter_rows(min_row=header_idx + 1, max_col=width or 1, values_only=True)
    raw_headers = next(records, ())
    headers = tuple(raw_headers[i] if i < len(raw_headers) and raw_headers[i] is not None
                    else f"Unnamed: {i}" for i in indices)
    projected = (tuple(row[i] if i < len(row) else None for i in indices) for row in records)
    return headers, projected


def _suggest_last_column(excel_path, sheet=None, header_idx=0, last_column=None):
    """اقتراح آخر عمود للفهرسة: آخر عمود فيه خليتان غير فارغتان على الأقل.
    فحص إضافي خفيف يُستخدم فقط عند تجاوز حد الأعمدة، وأي فشل فيه لا يغيّر رسالة الخطأ."""
    from openpyxl import load_workbook
    try:
        workbook = load_workbook(excel_path, read_only=True, data_only=True)
    except Exception:
        return None
    try:
        ws = workbook[sheet] if isinstance(sheet, str) else workbook.worksheets[sheet or 0]
        ws.reset_dimensions()
        counts = {}
        for row in ws.iter_rows(min_row=header_idx + 1, max_col=last_column, values_only=True):
            for index, value in enumerate(row):
                if value is not None and str(value).strip():
                    counts[index] = counts.get(index, 0) + 1
        dense = [index for index, count in counts.items() if count >= 2]
        return (max(dense) + 1) if dense else None
    except Exception:
        return None
    finally:
        workbook.close()


def _load_excel(excel_path, sheet=None, header_row=None, streaming_import=False, compact_columns=False, last_column=None):
    if header_row is None:
        header_idx = detect_header_row(excel_path, sheet)
    else:
        header_idx = max(0, int(header_row) - 1)  # المستخدم بيكتب رقم الصف زي ما شايفه في إكسل (1-indexed)

    if streaming_import:
        from openpyxl import load_workbook
        wb = load_workbook(excel_path, read_only=True, data_only=True)
        try:
            ws = wb[sheet] if isinstance(sheet, str) else wb.worksheets[sheet or 0]
            headers, records = _worksheet_records(ws, header_idx, compact_columns, last_column)
            names = [str(v) if v is not None else f"Unnamed: {i}" for i, v in enumerate(headers)]
            def values():
                for record in records:
                    row = ["" if v is None else str(v) for v in record[:len(names)]]
                    row.extend([""] * (len(names) - len(row)))
                    if any(v.strip() for v in row):
                        yield row
            df = pd.DataFrame(values(), columns=names)
        finally:
            wb.close()
    else:
        df = pd.read_excel(excel_path, sheet_name=sheet or 0, header=header_idx, dtype=str)
    df = df.fillna("")

    # شيل الأعمدة الفاضية بالكامل (زي أعمدة فاصلة بين الجداول) وأعمدة Unnamed
    # اللي مالهاش قيمة ولا حتى عنوان حقيقي
    if not streaming_import:
        empty_cols = [c for c in df.columns if (df[c].astype(str).str.strip() == "").all()]
        df = df.drop(columns=empty_cols)
    else:
        df = df.loc[df.apply(lambda row: any(str(v).strip() for v in row), axis=1)].reset_index(drop=True)

    original_columns = list(df.columns)
    columns, used = [], {"id", "_hash", "_key", "_overrides", "rank"}
    for c in df.columns:
        base = sanitize_col_name(c)
        clean, suffix = base, 2
        while streaming_import and clean.casefold() in used:
            clean = f"{base}_{suffix}"
            suffix += 1
        used.add(clean.casefold())
        columns.append(clean)
    df.columns = columns
    return df, original_columns, header_idx


def _rebuild_search_indexes(conn, search_columns):
    """يعيد بناء فهارس البحث (FTS5 + النسخة المطبّعة) بالكامل من جدول data الحالي.
    عملية رخيصة نسبياً حتى لعشرات آلاف الصفوف، وتضمن اتساق الفهرس دايماً."""
    cur = conn.cursor()
    fts_cols = ", ".join(f'"{c}"' for c in search_columns)

    cur.execute("DROP TABLE IF EXISTS data_fts")
    cur.execute(f"""
        CREATE VIRTUAL TABLE data_fts USING fts5(
            {fts_cols}, content='data', content_rowid='id',
            tokenize='unicode61 remove_diacritics 2'
        )
    """)
    cur.execute(f"INSERT INTO data_fts (rowid, {fts_cols}) SELECT id, {fts_cols} FROM data")

    cur.execute("DROP TABLE IF EXISTS data_normalized")
    cur.execute("CREATE VIRTUAL TABLE data_normalized USING fts5(text_norm, content='', tokenize='unicode61')")
    reader = conn.execute(f"SELECT id, {fts_cols} FROM data")
    while rows := reader.fetchmany(250):
        cur.executemany(
            "INSERT INTO data_normalized (rowid, text_norm) VALUES (?, ?)",
            ((r[0], normalize_arabic(" ".join(str(v or "") for v in r[1:]))) for r in rows),
        )


# ----------------------------------------------------------------------
# البناء الكامل (أول مرة فقط، أو عند طلب إعادة بناء كاملة)
# ----------------------------------------------------------------------
def build_db(excel_path, db_path, sheet=None, text_cols=None, unique_col=None, header_row=None, progress=None, last_column=None):
    """Build in a temporary database; publish only after successful indexing."""
    if last_column is not None and str(last_column).strip():
        from openpyxl.utils.cell import column_index_from_string
        raw = str(last_column).strip().upper()
        try:
            last_column = int(raw) if raw.isdigit() else column_index_from_string(raw)
        except ValueError:
            raise ValueError("آخر عمود يجب أن يكون حرفًا مثل N أو رقمًا مثل 14.") from None
        if not 1 <= last_column <= 16384:
            raise ValueError("آخر عمود يجب أن يكون بين A وXFD (1 إلى 16384).")
    else:
        last_column = None
    target = Path(db_path).resolve()
    fd, staging = tempfile.mkstemp(prefix=".index-", suffix=".sqlite", dir=target.parent)
    os.close(fd)
    try:
        if Path(excel_path).suffix.lower() in (".xlsx", ".xlsm"):
            result = _build_streaming(excel_path, staging, sheet, text_cols, unique_col, header_row, progress, last_column)
        else:
            if last_column is not None:
                raise ValueError("اختيار آخر عمود متاح لملفات XLSX وXLSM فقط.")
            result = _build_dataframe(excel_path, staging, sheet, text_cols, unique_col, header_row, progress)
        os.replace(staging, target)
        return result
    finally:
        if os.path.exists(staging):
            os.remove(staging)


def _build_streaming(excel_path, db_path, sheet, text_cols, unique_col, header_row, progress, last_column=None):
    from openpyxl import load_workbook
    report = progress or (lambda message: None)
    report("قراءة ملف الإكسل على دفعات...")
    header_idx = detect_header_row(excel_path, sheet) if header_row is None else max(0, int(header_row) - 1)
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    conn = sqlite3.connect(db_path)
    try:
        ws = workbook[sheet] if isinstance(sheet, str) else workbook.worksheets[sheet or 0]
        # Ignore inflated worksheet dimensions caused by formatting empty cells.
        headers, records = _worksheet_records(ws, header_idx, last_column=last_column)
        if not headers:
            raise ValueError("ملف الإكسل لا يحتوي على عناوين.")
        originals = [str(v) if v is not None else f"Unnamed: {i}" for i, v in enumerate(headers)]
        columns, used = [], {"id", "_hash", "_key", "_overrides", "rank"}
        for original in originals:
            base = sanitize_col_name(original)
            clean, suffix = base, 2
            while clean.casefold() in used:
                clean = f"{base}_{suffix}"
                suffix += 1
            used.add(clean.casefold())
            columns.append(clean)
        uc = sanitize_col_name(unique_col) if unique_col else None
        if unique_col in originals:
            uc = columns[originals.index(unique_col)]
        if uc and uc not in columns:
            raise ValueError(f"العمود الفريد '{unique_col}' غير موجود في الملف.")
        column_limit = min(conn.getlimit(sqlite3.SQLITE_LIMIT_COLUMN) - 4,
                           conn.getlimit(sqlite3.SQLITE_LIMIT_VARIABLE_NUMBER) - 3)
        if len(columns) > column_limit:
            hint = ""
            suggestion = _suggest_last_column(excel_path, sheet, header_idx, last_column)
            if suggestion and suggestion <= column_limit:
                from openpyxl.utils import get_column_letter
                hint = (f" اقتراح تقديري حسب كثافة البيانات: "
                        f"{get_column_letter(suggestion)} ({suggestion}). راجع النطاق قبل اختياره.")
            raise ValueError(
                f"الشيت يحتوي على {len(columns)} عمودًا غير فارغ بعد استبعاد الأعمدة الفارغة، "
                f"والحد المتاح هو {column_limit}. حدد آخر عمود للجدول في خانة آخر عمود للفهرسة، أو قسّم الشيت إلى ملفات أقل أعمدة. "
                "لم يتم حذف أي بيانات." + hint
            )
        definitions = ', '.join(f'"{c}" TEXT' for c in columns)
        conn.execute(f'CREATE TABLE data (id INTEGER PRIMARY KEY, _hash TEXT, _key TEXT, _overrides TEXT DEFAULT \'[]\', {definitions})')
        conn.execute('CREATE ' + ('UNIQUE ' if uc else '') + 'INDEX idx_key ON data(_key)')
        names = ', '.join(f'"{c}"' for c in columns)
        placeholders = ', '.join('?' for _ in range(len(columns) + 3))
        sql = f'INSERT INTO data (id, _hash, _key, {names}) VALUES ({placeholders})'
        batch, count = [], 0
        for record in records:
            values = ["" if v is None else str(v) for v in record[:len(columns)]]
            values.extend([""] * (len(columns) - len(values)))
            if not any(v.strip() for v in values):
                continue
            count += 1
            key = values[columns.index(uc)] if uc else str(count)
            if uc and not key.strip():
                raise ValueError("عمود المفتاح الفريد يحتوي على خلية فارغة.")
            batch.append((count, _row_hash(values), key, *values))
            if len(batch) >= 250:
                conn.executemany(sql, batch)
                batch.clear()
                report(f"تمت قراءة {count} صف...")
        if batch:
            conn.executemany(sql, batch)
        wanted = {sanitize_col_name(c.strip()) for c in text_cols.split(',')} if text_cols else None
        search_columns = [c for c in columns if wanted is None or c in wanted]
        if not search_columns:
            raise ValueError("لا توجد أعمدة قابلة للبحث.")
        report("تجهيز فهارس البحث...")
        _rebuild_search_indexes(conn, search_columns)
        for key, value in {
            "excel_path": str(Path(excel_path).resolve()), "sheet": sheet,
            "header_row": header_idx + 1, "search_columns": search_columns,
            "unique_col": uc, "original_columns": originals, "clean_columns": columns,
            "file_hash": _file_hash(Path(excel_path)),
            "last_sync": time.strftime("%Y-%m-%d %H:%M:%S"),
            "last_sync_stats": {"added": count, "updated": 0, "removed": 0},
            "streaming_import": True,
            "compact_columns": True,
            "last_column": last_column,
        }.items():
            set_config(conn, key, value)
        conn.commit()
        return {"added": count, "updated": 0, "removed": 0, "unchanged": 0}
    except sqlite3.IntegrityError as exc:
        raise ValueError("عمود المفتاح الفريد يحتوي على قيم مكررة. اختار عمودًا آخر.") from exc
    finally:
        conn.close()
        workbook.close()


def _build_dataframe(excel_path, db_path, sheet=None, text_cols=None, unique_col=None, header_row=None, progress=None):
    excel_path = Path(excel_path)
    if not excel_path.exists():
        raise FileNotFoundError(f"الملف غير موجود: {excel_path}")

    def report(msg):
        if progress:
            progress(msg)

    report("📖 بقرأ ملف الإكسل...")
    df, original_columns, header_idx = _load_excel(excel_path, sheet, header_row)
    if header_row is None and header_idx > 0:
        report(f"📌 لقيت صف العناوين الحقيقي في صف رقم {header_idx + 1} (فيه بيانات/عنوان فوقه اتجاهل)")
    report(f"✅ عدد الصفوف: {len(df)} | عدد الأعمدة: {len(df.columns)}")

    if text_cols:
        wanted = {sanitize_col_name(c.strip()) for c in text_cols.split(",")}
        search_columns = [c for c in df.columns if c in wanted]
        if not search_columns:
            raise ValueError("الأعمدة المطلوبة في text_cols غير موجودة في الملف.")
    else:
        search_columns = list(df.columns)

    uc = sanitize_col_name(unique_col) if unique_col else None
    if uc and uc not in df.columns:
        raise ValueError(f"العمود الفريد '{unique_col}' غير موجود في الملف.")

    conn = sqlite3.connect(db_path)
    cur = conn.cursor()

    cur.execute("DROP TABLE IF EXISTS data")
    cols_def = ", ".join(f'"{c}" TEXT' for c in df.columns)
    cur.execute(f'CREATE TABLE data (id INTEGER PRIMARY KEY, "_hash" TEXT, "_key" TEXT, "_overrides" TEXT DEFAULT \'[]\', {cols_def})')

    df = df.copy()
    df.insert(0, "id", range(1, len(df) + 1))
    df["_hash"] = df[[c for c in df.columns if c != "id"]].apply(lambda r: _row_hash(r.values), axis=1)
    df["_key"] = df[uc].astype(str) if uc else df["id"].astype(str)

    report("💾 بيتكتب في القاعدة...")
    df.to_sql("data", conn, if_exists="append", index=False)

    if uc:
        cur.execute('CREATE UNIQUE INDEX IF NOT EXISTS idx_key ON data("_key")')
    else:
        cur.execute('CREATE INDEX IF NOT EXISTS idx_key ON data("_key")')

    report("🔎 بيجهز فهرس البحث النصي...")
    _rebuild_search_indexes(conn, search_columns)

    set_config(conn, "excel_path", str(excel_path.resolve()))
    set_config(conn, "sheet", sheet)
    set_config(conn, "header_row", header_idx + 1)
    set_config(conn, "search_columns", search_columns)
    set_config(conn, "unique_col", uc)
    set_config(conn, "original_columns", original_columns)
    set_config(conn, "clean_columns", list(df.columns))
    set_config(conn, "file_hash", _file_hash(excel_path))
    set_config(conn, "last_sync", time.strftime("%Y-%m-%d %H:%M:%S"))
    set_config(conn, "last_sync_stats", {"added": len(df), "updated": 0, "removed": 0})

    conn.commit()
    conn.close()
    report(f"🎉 تم بناء القاعدة بنجاح: {len(df)} صف")
    return {"added": len(df), "updated": 0, "removed": 0, "unchanged": 0}


# ----------------------------------------------------------------------
# المزامنة التزايدية (Incremental Sync)
# ----------------------------------------------------------------------
def sync_db(db_path, excel_path=None, progress=None, force=False):
    """يقارن ملف الإكسل الحالي بالقاعدة، ويحدّث فقط الصفوف الجديدة/المتغيّرة/المحذوفة.
    لازم يكون فيه build_db اتعمل قبل كده مرة واحدة على الأقل (عشان الإعدادات محفوظة)."""

    probe = sqlite3.connect(db_path)
    try:
        streaming = get_config(probe, "streaming_import", False) and get_config(probe, "compact_columns", False)
    finally:
        probe.close()
    if streaming:
        return _sync_streaming(db_path, excel_path, progress, force)

    def report(msg):
        if progress:
            progress(msg)

    conn = sqlite3.connect(db_path)
    _ensure_overrides_column(conn)
    cfg_excel_path = excel_path or get_config(conn, "excel_path")
    if not cfg_excel_path:
        conn.close()
        raise ValueError("لا يوجد ملف إكسل مرتبط بالقاعدة. استخدم build أولاً.")

    excel_path = Path(cfg_excel_path)
    if not excel_path.exists():
        conn.close()
        raise FileNotFoundError(f"الملف غير موجود: {excel_path}")

    new_hash = _file_hash(excel_path)
    old_hash = get_config(conn, "file_hash")
    if not force and new_hash == old_hash:
        conn.close()
        report("✅ الملف لم يتغيّر منذ آخر مزامنة — مفيش حاجة تتحدّث.")
        return {"added": 0, "updated": 0, "removed": 0, "unchanged": None, "skipped": True}

    sheet = get_config(conn, "sheet")
    header_row = get_config(conn, "header_row")
    search_columns = get_config(conn, "search_columns")
    uc = get_config(conn, "unique_col")

    report("📖 بيقرأ الملف الجديد...")
    df, _, _ = _load_excel(excel_path, sheet, header_row,
                           get_config(conn, "streaming_import", False),
                           get_config(conn, "compact_columns", False),
                           get_config(conn, "last_column"))
    expected_columns = [c for c in _existing_columns(conn)
                        if c not in ("id", "_hash", "_key", "_overrides")
                        and c not in (get_config(conn, "local_columns", []) or [])]
    if list(df.columns) != expected_columns:
        conn.close()
        raise ValueError("عناوين الأعمدة تغيّرت. أضف الملف كقاعدة جديدة للحفاظ على التعديلات المحلية.")
    if uc and (uc not in df.columns or df[uc].eq("").any() or df[uc].duplicated().any()):
        conn.close()
        raise ValueError("عمود المفتاح الفريد مفقود أو يحتوي على قيم فارغة أو مكررة.")
    df["_key"] = df[uc].astype(str) if uc and uc in df.columns else (df.index + 1).astype(str)
    value_cols = [c for c in df.columns if c not in ("_key",)]
    df["_hash"] = df[value_cols].apply(lambda r: _row_hash(r.values), axis=1)

    if not uc and (get_config(conn, 'local_columns', []) or conn.execute("SELECT 1 FROM data WHERE _overrides IS NOT NULL AND _overrides != '[]' LIMIT 1").fetchone()):
        current = dict(conn.execute('SELECT _key,_hash FROM data'))
        incoming = dict(zip(df['_key'],df['_hash']))
        if current != incoming:
            conn.close()
            raise ValueError('أُوقفت المزامنة لحماية التعديلات المحلية: المصدر تغيّر دون مفتاح فريد يثبت هوية الصفوف.')

    cur = conn.cursor()
    existing = {row[0]: (row[1], row[2]) for row in cur.execute('SELECT "_key", id, "_hash" FROM data')}

    new_keys = set(df["_key"])
    old_keys = set(existing.keys())

    to_insert = df[df["_key"].apply(lambda k: k not in existing)]
    to_check = df[df["_key"].apply(lambda k: k in existing)]
    to_delete_keys = old_keys - new_keys

    updated_count = 0
    protected_count = 0
    if len(to_check):
        for _, row in to_check.iterrows():
            key = row["_key"]
            old_id, old_hash_val = existing[key]
            if row["_hash"] != old_hash_val:
                overridden = _get_overrides(conn, old_id)
                cols_to_update = [c for c in df.columns if c not in overridden]
                if overridden:
                    protected_count += 1
                if cols_to_update:
                    set_clause = ", ".join(f'"{c}"=?' for c in cols_to_update)
                    cur.execute(
                        f'UPDATE data SET {set_clause} WHERE id=?',
                        [*(row[c] for c in cols_to_update), old_id],
                    )
                updated_count += 1

    if len(to_insert):
        max_id = cur.execute("SELECT COALESCE(MAX(id),0) FROM data").fetchone()[0]
        to_insert = to_insert.copy()
        to_insert.insert(0, "id", range(max_id + 1, max_id + 1 + len(to_insert)))
        to_insert.to_sql("data", conn, if_exists="append", index=False)

    removed_count = 0
    if to_delete_keys:
        ids_to_remove = [existing[k][0] for k in to_delete_keys]
        cur.executemany("DELETE FROM data WHERE id=?", [(i,) for i in ids_to_remove])
        removed_count = len(ids_to_remove)

    report("🔎 بيعيد بناء فهرس البحث...")
    _rebuild_search_indexes(conn, search_columns)

    set_config(conn, "file_hash", new_hash)
    set_config(conn, "last_sync", time.strftime("%Y-%m-%d %H:%M:%S"))
    stats = {
        "added": int(len(to_insert)),
        "updated": int(updated_count),
        "removed": int(removed_count),
        "unchanged": int(len(to_check) - updated_count),
        "protected": int(protected_count),
        "skipped": False,
    }
    set_config(conn, "last_sync_stats", stats)

    conn.commit()
    conn.close()
    report(f"🎉 المزامنة خلصت: {stats['added']} جديد، {stats['updated']} اتحدّث، {stats['removed']} اتشال"
           + (f"، {stats['protected']} صف فيه خلايا محمية اتصانت" if protected_count else ""))
    return stats


def _sync_streaming(db_path, excel_path=None, progress=None, force=False):
    """Stage the source on disk, then apply changes in one rollback-safe transaction."""
    report = progress or (lambda _: None)
    conn = sqlite3.connect(db_path)
    try:
        source = Path(excel_path or get_config(conn, 'excel_path'))
        digest = _file_hash(source)
        if not force and digest == get_config(conn, 'file_hash'):
            return {'added':0,'updated':0,'removed':0,'unchanged':None,'skipped':True}
        uc = get_config(conn, 'unique_col')
        with tempfile.TemporaryDirectory(prefix='excel-sync-') as temp:
            staging = str(Path(temp)/'incoming.sqlite')
            build_db(source, staging, sheet=get_config(conn,'sheet'), unique_col=uc,
                     header_row=get_config(conn,'header_row'), last_column=get_config(conn,'last_column'),
                     progress=report)
            conn.execute('ATTACH DATABASE ? AS incoming', (staging,))
            reader = None
            try:
                conn.execute('BEGIN IMMEDIATE')
                _ensure_overrides_column(conn)
                local = get_config(conn,'local_columns',[]) or []
                columns = [c for c in _existing_columns(conn) if c not in ('id','_key','_hash','_overrides') and c not in local]
                incoming_columns = [r[1] for r in conn.execute('PRAGMA incoming.table_info(data)') if r[1] not in ('id','_key','_hash','_overrides')]
                if columns != incoming_columns:
                    raise ValueError('عناوين الأعمدة تغيّرت. أضف الملف كقاعدة جديدة للحفاظ على التعديلات المحلية.')
                if not uc and (local or conn.execute("SELECT 1 FROM data WHERE _overrides IS NOT NULL AND _overrides != '[]' LIMIT 1").fetchone()):
                    changed = conn.execute('SELECT 1 FROM incoming.data n LEFT JOIN main.data d ON n._key=d._key WHERE d.id IS NULL OR n._hash IS NOT d._hash LIMIT 1').fetchone()
                    deleted = conn.execute('SELECT 1 FROM main.data d WHERE NOT EXISTS(SELECT 1 FROM incoming.data n WHERE n._key=d._key) LIMIT 1').fetchone()
                    if changed or deleted:
                        raise ValueError('أُوقفت المزامنة لحماية تعديلاتك: المصدر تغيّر ولا يوجد مفتاح فريد يثبت هوية الصفوف. أضف المصدر كملف مستقل بمفتاح فريد قبل نقل التعديلات.')
                stats = {'added':0,'updated':0,'removed':0,'unchanged':0,'protected':0,'skipped':False}
                next_id = conn.execute('SELECT COALESCE(MAX(id),0)+1 FROM main.data').fetchone()[0]
                names = ','.join('"'+c.replace('"','""')+'"' for c in columns)
                reader = conn.execute(f'SELECT n.id,n._key,n._hash,{",".join("n."+chr(34)+c+chr(34) for c in columns)},d.id,d._hash,d._overrides FROM incoming.data n LEFT JOIN main.data d ON n._key=d._key ORDER BY n.id')
                for number, row in enumerate(reader,1):
                    if number % 250 == 0:
                        report(f'مزامنة {number} صف...')
                    _, key, row_hash = row[:3]
                    values = row[3:3+len(columns)]
                    old_id, old_hash, overrides = row[-3:]
                    if old_id is None:
                        placeholders = ','.join('?' for _ in range(len(columns)+3))
                        conn.execute(f'INSERT INTO main.data(id,_key,_hash,{names}) VALUES({placeholders})', (next_id,key,row_hash,*values))
                        next_id += 1
                        stats['added'] += 1
                    elif old_hash != row_hash:
                        protected = set(json.loads(overrides or '[]'))
                        pairs = [(c,v) for c,v in zip(columns,values) if c not in protected]
                        assignments = ','.join('"'+c+'"=?' for c,_ in pairs)
                        conn.execute('UPDATE main.data SET '+(assignments+',' if assignments else '')+'_hash=? WHERE id=?', [*(v for _,v in pairs),row_hash,old_id])
                        stats['updated'] += 1
                        stats['protected'] += bool(protected)
                    else:
                        stats['unchanged'] += 1
                stats['removed'] = conn.execute('DELETE FROM main.data WHERE NOT EXISTS(SELECT 1 FROM incoming.data n WHERE n._key=main.data._key)').rowcount
                report('تحديث فهارس البحث...')
                _rebuild_search_indexes(conn, get_config(conn,'search_columns'))
                report('حفظ المزامنة...')
                set_config(conn,'file_hash',digest)
                set_config(conn,'excel_path',str(source.resolve()))
                set_config(conn,'last_sync',time.strftime('%Y-%m-%d %H:%M:%S'))
                set_config(conn,'last_sync_stats',stats)
                conn.commit()
                return stats
            except BaseException:
                conn.rollback()
                raise
            finally:
                if reader is not None:
                    reader.close()
                conn.execute('DETACH DATABASE incoming')
    finally:
        conn.close()


# ----------------------------------------------------------------------
# البحث
# ----------------------------------------------------------------------
def _escape_fts_term(term: str) -> str:
    """يحوّل كلمة بحث لصيغة 'phrase' آمنة داخل FTS5، عشان رموز زي - و : و ^
    ما تتفسّرش كأوامر بحث خاصة (NOT / column-filter) بدل ما تتقرأ كنص حرفي."""
    escaped = term.replace('"', '""')
    return f'"{escaped}"'


def _build_fts_query(raw_query: str) -> str:
    """يبني عبارة MATCH آمنة: كل كلمة مفصولة بمسافة بتتحول لعبارة حرفية (phrase)،
    والمسافة بين الكلمات لسه بتعني AND عادي. كده أي كود زي BP-BW-BL-04 أو أرقام
    فيها شرطات هيتدوّر عليه كنص كامل بدل ما يتفسّر كأوامر استبعاد."""
    terms = raw_query.strip().split()
    if not terms:
        return raw_query
    return " ".join(_escape_fts_term(t) for t in terms)


def search(db_path, query, limit=20, offset=0):
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    cur = conn.cursor()

    fts_query = _build_fts_query(query)
    rows = []
    total = 0
    try:
        rows = cur.execute(
            """
            SELECT data.*, bm25(data_fts) AS rank FROM data
            JOIN data_fts ON data.id = data_fts.rowid
            WHERE data_fts MATCH ?
            ORDER BY rank LIMIT ? OFFSET ?
            """,
            (fts_query, limit, offset),
        ).fetchall()
        total = cur.execute(
            "SELECT COUNT(*) FROM data JOIN data_fts ON data.id = data_fts.rowid WHERE data_fts MATCH ?",
            (fts_query,),
        ).fetchone()[0]
    except sqlite3.OperationalError:
        rows = []

    if total == 0:
        norm_query = _build_fts_query(normalize_arabic(query))
        try:
            rows = cur.execute(
                """
                SELECT data.* FROM data
                JOIN data_normalized ON data.id = data_normalized.rowid
                WHERE data_normalized.text_norm MATCH ?
                LIMIT ? OFFSET ?
                """,
                (norm_query, limit, offset),
            ).fetchall()
            total = cur.execute(
                "SELECT COUNT(*) FROM data JOIN data_normalized ON data.id = data_normalized.rowid "
                "WHERE data_normalized.text_norm MATCH ?",
                (norm_query,),
            ).fetchone()[0]
        except sqlite3.OperationalError:
            rows, total = [], 0

    conn.close()
    result = []
    for r in rows:
        d = dict(r)
        d.pop("_hash", None)
        d.pop("_key", None)
        d.pop("rank", None)
        overrides_raw = d.pop("_overrides", None)
        try:
            d["_protected"] = json.loads(overrides_raw) if overrides_raw else []
        except (TypeError, json.JSONDecodeError):
            d["_protected"] = []
        result.append(d)
    return result, total


def query_rows(db_path, query="", filters=None, limit=20, offset=0):
    """Combine validated column filters with the existing full-text search."""
    if not isinstance(filters, (list, type(None))) or len(filters or []) > 12:
        raise ValueError("الفلاتر يجب أن تكون قائمة من 12 شرطًا كحد أقصى.")
    if query and not filters:
        return search(db_path, query, limit, offset)
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    conn.create_function("norm", 1, lambda value: normalize_arabic(value).casefold())
    try:
        allowed = set(get_config(conn, "search_columns", []) or [])
        conditions, values = [], []
        for item in filters or []:
            from row_filters import predicate
            clause, params = predicate(item, allowed, lambda value: normalize_arabic(value).casefold())
            conditions.append(clause)
            values.extend(params)
        base = " AND ".join(conditions) or "1=1"
        candidates = [(base, values)]
        if query:
            candidates = [
                (base + " AND id IN (SELECT rowid FROM data_fts WHERE data_fts MATCH ?)", values + [_build_fts_query(query)]),
                (base + " AND id IN (SELECT rowid FROM data_normalized WHERE text_norm MATCH ?)", values + [_build_fts_query(normalize_arabic(query))]),
            ]
        for where, params in candidates:
            try:
                total = conn.execute(f"SELECT COUNT(*) FROM data WHERE {where}", params).fetchone()[0]
                if not total:
                    continue
                rows = conn.execute(f"SELECT * FROM data WHERE {where} ORDER BY id LIMIT ? OFFSET ?",
                                    [*params, limit, offset]).fetchall()
            except sqlite3.OperationalError:
                if not query:
                    raise
                continue
            result = []
            for row in rows:
                item = dict(row)
                item.pop('_hash', None)
                item.pop('_key', None)
                raw = item.pop('_overrides', None)
                try:
                    item['_protected'] = json.loads(raw) if raw else []
                except (TypeError, json.JSONDecodeError):
                    item['_protected'] = []
                result.append(item)
            return result, total
        return [], 0
    finally:
        conn.close()


def get_stats(db_path):
    conn = sqlite3.connect(db_path)
    try:
        count = conn.execute("SELECT COUNT(*) FROM data").fetchone()[0]
    except sqlite3.OperationalError:
        count = 0
    cfg = {
        "excel_path": get_config(conn, "excel_path"),
        "last_sync": get_config(conn, "last_sync"),
        "last_sync_stats": get_config(conn, "last_sync_stats"),
        "search_columns": get_config(conn, "search_columns"),
        "original_columns": get_config(conn, "original_columns"),
        "unique_col": get_config(conn, "unique_col"),
    }
    conn.close()
    cfg["row_count"] = count
    return cfg


# ----------------------------------------------------------------------
# مقارنة ملفين (Diff) — بيحتاج عمود مفتاح مشترك المعنى في الملفين
# (ممكن يكون بأسماء مختلفة، مثلاً "الكود" في ملف و"رمز_البند" في التاني)
# ----------------------------------------------------------------------
def get_columns(db_path):
    stats = get_stats(db_path)
    return stats.get("search_columns") or []


def _norm_key(v) -> str:
    return normalize_arabic(str(v).strip())


def _fetch_keyed_rows(db_path, key_col):
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    cols = [r[1] for r in conn.execute("PRAGMA table_info(data)").fetchall()
            if r[1] not in ("id", "_hash", "_key", "_overrides")]
    if key_col not in cols:
        conn.close()
        raise ValueError(f"العمود '{key_col}' مش موجود في هذه القاعدة.")
    rows = conn.execute("SELECT * FROM data").fetchall()
    conn.close()
    out = {}
    for r in rows:
        d = {k: r[k] for k in cols}
        raw_key = str(d.get(key_col, "")).strip()
        if not raw_key:
            continue
        normalized_key = _norm_key(raw_key)
        if normalized_key in out:
            raise ValueError(f"عمود المقارنة '{key_col}' فيه مفتاح مكرر: {raw_key}. اختار عمودًا فريدًا لتجنب فقد صفوف من المقارنة.")
        out[normalized_key] = {"raw_key": raw_key, "row": d}
    return out, cols


def compare_databases(db_path_a, db_path_b, key_col_a, key_col_b, query=""):
    rows_a, cols_a = _fetch_keyed_rows(db_path_a, key_col_a)
    rows_b, cols_b = _fetch_keyed_rows(db_path_b, key_col_b)

    terms = normalize_arabic(query).casefold().split()
    if terms:
        def matches(item):
            text = normalize_arabic(" ".join(str(v or "") for v in item["row"].values())).casefold()
            return all(term in text for term in terms)
        selected = {key for rows in (rows_a, rows_b) for key, item in rows.items() if matches(item)}
        rows_a = {key: item for key, item in rows_a.items() if key in selected}
        rows_b = {key: item for key, item in rows_b.items() if key in selected}

    shared_cols = [c for c in cols_a if c in cols_b and c not in (key_col_a, key_col_b)]

    keys_a = set(rows_a.keys())
    keys_b = set(rows_b.keys())
    common_keys = keys_a & keys_b
    only_a_keys = keys_a - keys_b
    only_b_keys = keys_b - keys_a

    changed = []
    identical_count = 0
    for k in common_keys:
        row_a = rows_a[k]["row"]
        row_b = rows_b[k]["row"]
        diffs = {}
        for c in shared_cols:
            va, vb = str(row_a.get(c, "")).strip(), str(row_b.get(c, "")).strip()
            if va != vb:
                diffs[c] = {"a": va, "b": vb}
        if diffs:
            changed.append({
                "key": rows_a[k]["raw_key"],
                "diffs": diffs,
                "row_a": row_a,
                "row_b": row_b,
            })
        else:
            identical_count += 1

    only_in_a = [rows_a[k]["row"] for k in only_a_keys]
    only_in_b = [rows_b[k]["row"] for k in only_b_keys]

    return {
        "query": query,
        "shared_columns": shared_cols,
        "matched_identical": identical_count,
        "matched_changed": len(changed),
        "only_in_a_count": len(only_in_a),
        "only_in_b_count": len(only_in_b),
        "changed": changed,
        "only_in_a": only_in_a,
        "only_in_b": only_in_b,
    }


def search_multi(db_paths_and_names, query, limit_per_file=10):
    """بحث نفس الكلمة في أكتر من قاعدة مرة واحدة، بيرجّع النتائج مجمّعة حسب الملف."""
    results = []
    for db_id, name, db_path in db_paths_and_names:
        if not os.path.exists(db_path):
            continue
        rows, total = search(db_path, query, limit=limit_per_file, offset=0)
        if total:
            results.append({"db_id": db_id, "name": name, "rows": rows, "total": total})
    return results


# ----------------------------------------------------------------------
# تعديل الخلايا وإضافة أعمدة محلية (بتفضل زي ما هي مهما اتعملت مزامنة)
# ----------------------------------------------------------------------
def _existing_columns(conn):
    return [r[1] for r in conn.execute("PRAGMA table_info(data)").fetchall()]


def _ensure_overrides_column(conn):
    """توافق مع قواعد قديمة اتبنت قبل ميزة (حماية الخلايا المعدّلة يدوياً)."""
    if "_overrides" not in _existing_columns(conn):
        conn.execute('ALTER TABLE data ADD COLUMN "_overrides" TEXT DEFAULT \'[]\'')
        conn.execute("UPDATE data SET \"_overrides\" = '[]' WHERE \"_overrides\" IS NULL")


def _get_overrides(conn, row_id):
    row = conn.execute('SELECT "_overrides" FROM data WHERE id=?', (row_id,)).fetchone()
    if not row or not row[0]:
        return set()
    try:
        return set(json.loads(row[0]))
    except (TypeError, json.JSONDecodeError):
        return set()


def add_local_column(db_path, col_name, default_value=""):
    conn = sqlite3.connect(db_path)
    clean = sanitize_col_name(col_name)
    cols = _existing_columns(conn)
    if clean in cols:
        conn.close()
        raise ValueError(f"العمود '{col_name}' موجود بالفعل.")

    conn.execute(f'ALTER TABLE data ADD COLUMN "{clean}" TEXT')
    conn.execute(f'UPDATE data SET "{clean}" = ?', (default_value or "",))

    local_cols = get_config(conn, "local_columns", []) or []
    local_cols.append(clean)
    set_config(conn, "local_columns", local_cols)

    search_cols = get_config(conn, "search_columns", []) or []
    if clean not in search_cols:
        search_cols.append(clean)
        set_config(conn, "search_columns", search_cols)
    _rebuild_search_indexes(conn, search_cols)

    orig_cols = get_config(conn, "original_columns", []) or []
    if col_name not in orig_cols:
        orig_cols.append(col_name)
        set_config(conn, "original_columns", orig_cols)

    conn.commit()
    conn.close()
    return clean


def update_cell(db_path, row_id, col_name, value):
    conn = sqlite3.connect(db_path)
    _ensure_overrides_column(conn)
    cols = _existing_columns(conn)
    if col_name not in cols or col_name in ("id", "_hash", "_key", "_overrides"):
        conn.close()
        raise ValueError("عمود غير قابل للتعديل.")

    conn.execute(f'UPDATE data SET "{col_name}" = ? WHERE id = ?', (value, row_id))

    # لو العمود ده أصله من الإكسل (مش عمود محلي)، سجّله كـ"محمي" عشان المزامنة
    # الجاية متلمسوش وترجّعه لقيمة الإكسل القديمة. الأعمدة المحلية أصلاً محمية
    # دايماً بدون الحاجة لده (المزامنة مبتلمسهاش خالص).
    local_cols = set(get_config(conn, "local_columns", []) or [])
    if col_name not in local_cols:
        overrides = _get_overrides(conn, row_id)
        overrides.add(col_name)
        conn.execute('UPDATE data SET "_overrides" = ? WHERE id = ?', (json.dumps(list(overrides)), row_id))

    conn.commit()

    search_cols = get_config(conn, "search_columns", []) or []
    if col_name in search_cols:
        _rebuild_search_indexes(conn, search_cols)
        conn.commit()
    conn.close()


def clear_override(db_path, row_id, col_name):
    """يلغي حماية خلية معدّلة يدوياً، فترجع تقبل قيمتها من الإكسل في أول مزامنة جاية."""
    conn = sqlite3.connect(db_path)
    _ensure_overrides_column(conn)
    overrides = _get_overrides(conn, row_id)
    overrides.discard(col_name)
    conn.execute('UPDATE data SET "_overrides" = ? WHERE id = ?', (json.dumps(list(overrides)), row_id))
    # A cleared override must be restored even if the Excel file is unchanged.
    conn.execute('UPDATE data SET "_hash" = NULL WHERE id = ?', (row_id,))
    set_config(conn, "file_hash", None)
    conn.commit()
    conn.close()


def is_local_column(db_path, col_name):
    conn = sqlite3.connect(db_path)
    local_cols = get_config(conn, "local_columns", []) or []
    conn.close()
    return col_name in local_cols
