"""Bounded-memory search exports, including multiple files and filter scope."""
import sqlite3
import math
import re
from datetime import datetime, timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
import excel_db_core as core
from row_filters import validate, missing_column


INK = '19354A'
GOLD = 'A97C2F'
ARABIC = re.compile(r'[\u0600-\u06ff]')
OPS = {'contains':'يحتوي على', 'not_contains':'لا يحتوي على',
       'equals':'يساوي', 'empty':'فارغ', 'not_empty':'غير فارغ'}


def _widths(columns, rows):
    widths = []
    for col in columns:
        longest = max([len(str(col))] + [max(map(len, str(row.get(col, '') or '').split('\n'))) for row in rows])
        description = any(word in col.lower() for word in ('وصف', 'بيان', 'ملاحظ', 'description', 'notes'))
        widths.append(68 if description or longest > 65 else min(38, max(14, longest + 3)))
    return widths


def _setup(sheet, widths, label):
    sheet.sheet_view.rightToLeft = True
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = GOLD
    sheet.freeze_panes = 'B2'
    sheet.print_title_rows = '1:1'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = '8' if len(widths) > 8 else '9'  # A3 / A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins = PageMargins(left=.3, right=.3, top=.55, bottom=.55, header=.2, footer=.2)
    # Escape Excel's header/footer control character in user-provided names.
    sheet.oddHeader.center.text = str(label).replace('&', '&&')[:120]
    sheet.oddFooter.center.text = 'صفحة &P من &N'
    sheet.oddFooter.right.text = 'ExcelSearch'
    for i, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(i)].width = width


def _append(sheet, values, widths, row_number, header=False):
    cells = []
    height = 32 if header else 25
    for i, value in enumerate(values):
        cell = WriteOnlyCell(sheet, value=value)
        if isinstance(value, str):
            cell.data_type = 's'  # Preserve leading zeroes and prevent formula injection.
        text = '' if value is None else str(value)
        rtl = bool(ARABIC.search(text))
        cell.font = Font(name='Arial', size=11, bold=header, color='FFFFFF' if header else INK)
        cell.fill = PatternFill('solid', fgColor=INK if header else ('F0F4F7' if row_number % 2 == 0 else 'FFFFFF'))
        cell.alignment = Alignment(horizontal='right' if rtl else 'left', vertical='top',
                                   wrap_text=True, readingOrder=2 if rtl else 1)
        cell.border = Border(bottom=Side(style='thin' if header else 'hair', color=GOLD if header else 'DCE4EA'))
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            cell.number_format = '#,##0.##########'
        lines = sum(max(1, math.ceil(len(part) / max(8, widths[i] - 4))) for part in text.split('\n'))
        height = max(height, min(409, lines * 16 + 8))
        cells.append(cell)
    sheet.row_dimensions[row_number].height = height
    sheet.append(cells)


def _finish(sheet, columns, last_row):
    end = get_column_letter(len(columns))
    sheet.auto_filter.ref = f'A1:{end}{last_row}'
    sheet.print_area = f'A1:{end}{last_row}'


def write_search_export(entries, query, filters, output):
    if not isinstance(filters, list) or len(filters) > 12:
        raise ValueError('قائمة فلاتر غير صالحة.')
    for item in filters:
        validate(item)
    wb=Workbook(write_only=True)
    wb.properties.creator='ExcelSearch'
    wb.properties.title='نتائج البحث والفلترة'
    wb.properties.language='ar-SA'
    summary=wb.create_sheet('الملفات')
    summary_widths=[38,24,20,70]
    _setup(summary,summary_widths,'ملخص نتائج البحث والفلترة')
    summary_line=0
    def append(sheet,values):
        nonlocal summary_line
        summary_line+=1
        _append(sheet,values,summary_widths,summary_line,header=summary_line==1)
    append(summary,['الملف','ورقة النتائج','عدد النتائج','الحالة'])
    for number,entry in enumerate(entries,1):
        path=entry['db_path']
        allowed=set(core.get_columns(path))
        if any(missing_column(item, allowed) for item in filters):
            append(summary,[entry.get('name',''),'',0,'مستبعد: أحد أعمدة الفلترة غير موجود'])
            continue
        rows,total=core.query_rows(path,query,filters,limit=500)
        conn=sqlite3.connect(path)
        try:
            columns=[row[1] for row in conn.execute('PRAGMA table_info(data)') if row[1] not in ('_hash','_key','_overrides')]
        finally:
            conn.close()
        title=f'النتائج {number}'
        widths=_widths(columns,rows)
        sheet=wb.create_sheet(title)
        _setup(sheet,widths,entry.get('name',title))
        _append(sheet,columns,widths,1,header=True)
        append(summary,[entry.get('name',''),title,total,'تم'])
        offset,line,part=0,1,1
        while rows:
            for row in rows:
                if line>=1048576:
                    _finish(sheet,columns,line)
                    part+=1;sheet=wb.create_sheet(f'{title}-{part}');line=1
                    _setup(sheet,widths,entry.get('name',title))
                    _append(sheet,columns,widths,1,header=True)
                    append(summary,[entry.get('name',''),sheet.title,'','تكملة النتائج'])
                line+=1
                _append(sheet,[row.get(col,'') for col in columns],widths,line)
            offset+=len(rows)
            if offset>=total:
                break
            rows,_=core.query_rows(path,query,filters,limit=500,offset=offset)
        _finish(sheet,columns,line)
    summary.auto_filter.ref=f'A1:D{summary_line}'
    append(summary,['','','',''])
    append(summary,['نطاق التقرير','القيمة','',''])
    append(summary,['كلمة البحث',query or 'كل النتائج','',''])
    append(summary,['وقت التصدير (UTC)',datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M'),'',''])
    for item in filters:
        append(summary,['فلتر: '+('كل الأعمدة' if item.get('scope')=='all' else item.get('col','')),
                        OPS[item.get('op','contains')],'',str(item.get('value',''))])
    if not filters:
        append(summary,['الفلاتر','بدون فلاتر إضافية','',''])
    summary.print_area=f'A1:D{summary_line}'
    wb.save(output)
