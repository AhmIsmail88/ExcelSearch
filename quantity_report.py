"""One result per row, with chosen context columns and year-by-year quantities."""
import json
import math
from openpyxl.utils import get_column_letter
from quantity_comparison import MODES


def write_quantity_report(engine, conn, wb, sheet, setup, append):
    context = [(tag, col) for tag in ('a', 'b') for col in engine.display_columns[tag]]
    headers = ['حالة المطابقة', 'التشابه النصي', 'المفتاح / الزوج', 'نتيجة الكميات']
    widths = [24, 16, 22, 46]
    for tag, col in context:
        headers.append(f'{col} | {engine.names[tag]}')
        widths.append(60 if 'وصف' in col or 'description' in col.lower() else 28)
    numeric_start = len(headers)
    for ca, cb in engine.qty_pairs:
        headers.extend([f'{ca} | الأول', f'{cb} | الثاني', f'الفرق: {ca} ← {cb}'])
        widths.extend([20, 20, 20])
    headers.extend(['إجمالي الأول', 'إجمالي الثاني', 'فرق الإجمالي'])
    widths.extend([22, 22, 22])
    last_col = get_column_letter(len(headers))
    formats = {index: '#,##0.######;[Red]-#,##0.######;0' for index in range(numeric_start, len(headers))}
    formats[1] = '0.0%'

    def start(target):
        setup(target, widths, 'مقارنة الكميات | بند واحد في كل صف',freeze=f'{get_column_letter(numeric_start+1)}6')
        target.row_dimensions[2].height = 38
        append(target, ['الملف الأول', engine.names['a'], 'الملف الثاني', engine.names['b']])
        target.row_dimensions[3].height = 45
        append(target, ['طريقة المقارنة', MODES[engine.quantity_mode], 'نطاق التقرير', 'الصفوف التي اجتازت الفلاتر فقط'])
        target.merged_cells.add(f'A4:{last_col}4')
        target.row_dimensions[4].height = 42
        append(target, ['الفرق = الثاني ناقص الأول. الفراغ محفوظ ولا يعني صفرًا؛ الإجمالي يجمع القيم الرقمية الموجودة فقط. كل اقتراح يحتاج مراجعة مستقلة ولا تجمع كميات الاقتراحات قبل قبول التطابق.'])
        target.row_dimensions[5].height = 66
        append(target, headers, fill='19354A', color='FFFFFF', bold=True)

    start(sheet)
    line, part = 5, 1
    from comparison_report import LABELS
    for section, score, raw in conn.execute('SELECT section,score,body FROM results ORDER BY id'):
        body = json.loads(raw); qty = body['qty']
        values = [LABELS[section], None if score is None else score/100, body['key'], qty['status']]
        values.extend(body['row_'+tag].get(col, '') for tag, col in context)
        highlights = set()
        for item in qty['cols']:
            if item['changed']:
                highlights.update([len(values), len(values)+1, len(values)+2])
            values.extend([item['a_num'] if item['a_num'] is not None else item['a'],
                           item['b_num'] if item['b_num'] is not None else item['b'], item['delta']])
        if qty['delta_total'] is not None and abs(qty['delta_total']) > 1e-6:
            highlights.add(len(values)+2)
        values.extend([qty['a_total'], qty['b_total'], qty['delta_total']])
        if line >= 1048576:
            sheet.auto_filter.ref = f'A5:{last_col}{line}'
            part += 1; sheet = wb.create_sheet(f'المقارنة {part}');start(sheet);line = 5
        line += 1
        height = max([3]+[sum(max(1, math.ceil(len(chunk)/max(12,widths[i]-5))) for chunk in str(v or '').split('\n')) for i,v in enumerate(values)])
        sheet.row_dimensions[line].height = min(390, height*16)
        append(sheet, values, fill='F6F8FA' if line%2==0 else 'FFFFFF', status=section,
               formats=formats, highlights=highlights)
    sheet.auto_filter.ref = f'A5:{last_col}{line}'
