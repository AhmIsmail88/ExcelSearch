import io
import json
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
import app
import excel_db_core as core


class APITests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.root=Path(self.temp.name)
        self.config=patch.object(app,'CONFIG_PATH',str(self.root/'config.json'))
        self.databases=patch.object(app,'DATABASES_DIR',str(self.root))
        self.config.start();self.databases.start()
        entries={}
        for name,description in [('a','توريد مضخة مياه'),('b','توريد مضخة مياه')]:
            wb=Workbook();wb.active.append(['code','description','qty'])
            wb.active.append(['1',description,'1' if name=='a' else '2'])
            path=self.root/(name+'.xlsx');wb.save(path);wb.close()
            db=self.root/(name+'.sqlite');core.build_db(path,db,unique_col='code',header_row=1)
            entries[name]={'name':name,'db_path':str(db),'excel_path':str(path)}
        app.save_app_config({'active_db_id':'b','databases':entries})
        self.client=app.app.test_client()

    def tearDown(self):
        app.cleanup_comparisons();app.comparisons.clear()
        self.databases.stop();self.config.stop();self.temp.cleanup()

    def test_mutation_requires_explicit_database_and_does_not_follow_active(self):
        payload={'id':1,'col':'qty','value':'99'}
        self.assertEqual(self.client.post('/api/update_cell',json=payload).status_code,400)
        self.assertEqual(self.client.post('/api/update_cell',json={**payload,'db_id':'missing'}).status_code,400)
        self.assertEqual(self.client.post('/api/update_cell',json={**payload,'db_id':'a'}).status_code,200)
        self.assertEqual(self.client.get('/api/search?db_id=a').json['rows'][0]['qty'],'99')
        self.assertEqual(self.client.get('/api/search?db_id=b').json['rows'][0]['qty'],'2')

    def test_async_comparison_pages_and_review_export(self):
        response=self.client.post('/api/compare',json={'db_id_a':'a','db_id_b':'b','key_col_a':'description','key_col_b':'description','mode':'similarity','async':True})
        self.assertEqual(response.status_code,202)
        token=response.json['job_id']
        for _ in range(100):
            job=self.client.get('/api/jobs/'+token).json
            if job['state'] in ('done','error'):
                break
            time.sleep(.05)
        self.assertEqual(job['state'],'done',job)
        result=job['result'];comparison_id=result['comparison_id']
        self.assertEqual(result['total'],1)
        self.assertEqual(self.client.post('/api/compare_decision',json={'comparison_id':comparison_id,'id':result['items'][0]['id'],'decision':'accepted'}).status_code,200)
        page=self.client.post('/api/compare_page',json={'comparison_id':comparison_id,'section':'accepted'}).json
        self.assertEqual(page['total'],1)
        qpage=self.client.post('/api/compare_page',json={'comparison_id':comparison_id,'qty_only':True}).json
        self.assertTrue(qpage['qty_only'])
        self.assertEqual(qpage['total'],0)
        export=self.client.post('/api/compare_export',json={'comparison_id':comparison_id})
        wb=load_workbook(io.BytesIO(export.data),read_only=True)
        self.assertEqual(list(wb['المقارنة'].values)[5][0],'مقبول يدويًا')
        wb.close();export.close()
        self.client.post('/api/compare_close',json={'comparison_id':comparison_id})
        self.assertEqual(self.client.post('/api/compare_page',json={'comparison_id':comparison_id}).json['section'],'accepted')
        self.client.post('/api/compare_delete',json={'comparison_id':comparison_id})
        self.assertEqual(self.client.post('/api/compare_page',json={'comparison_id':comparison_id}).status_code,400)

    def test_search_export_uses_literal_strings_and_filters(self):
        self.client.post('/api/update_cell',json={'db_id':'a','id':1,'col':'qty','value':'=1+1'})
        response=self.client.get('/api/export',query_string={'db_id':'a','filters':json.dumps([{'col':'qty','op':'equals','value':'=1+1'}])})
        self.assertEqual(response.status_code,200)
        wb=load_workbook(io.BytesIO(response.data))
        sheet=wb['النتائج 1']
        headers=[c.value for c in sheet[1]]
        cell=sheet.cell(2,headers.index('qty')+1)
        self.assertEqual((cell.value,cell.data_type),('=1+1','s'))
        wb.close();response.close()

    def test_all_file_filters_and_export(self):
        filters=json.dumps([{'col':'qty','op':'equals','value':'1'}])
        response=self.client.get('/api/search_all',query_string={'filters':filters})
        self.assertEqual(response.status_code,200)
        self.assertEqual([group['db_id'] for group in response.json['groups']],['a'])
        response=self.client.get('/api/export',query_string={'filters':filters,'all_files':'1'})
        wb=load_workbook(io.BytesIO(response.data))
        self.assertEqual(wb['النتائج 1'].max_row,2)
        self.assertEqual(wb['النتائج 2'].max_row,1)
        wb.close();response.close()

    def test_exclude_any_cell_search_export_and_comparison(self):
        # A match in a different column still excludes the entire row.
        self.client.post('/api/update_cell',json={'db_id':'a','id':1,'col':'qty','value':'خدمات المَسْجِد'})
        filters=[{'scope':'all','col':'','op':'not_contains','value':'المسجد'}]
        args={'filters':json.dumps(filters)}
        self.assertEqual(self.client.get('/api/search',query_string={**args,'db_id':'a'}).json['total'],0)
        self.assertEqual(self.client.get('/api/search',query_string={**args,'db_id':'b'}).json['total'],1)
        grouped=self.client.get('/api/search_all',query_string=args).json
        self.assertEqual([g['db_id'] for g in grouped['groups']],['b'])
        self.assertEqual(grouped['skipped'],[])
        response=self.client.get('/api/export',query_string={**args,'all_files':'1'})
        wb=load_workbook(io.BytesIO(response.data));self.assertEqual(wb['النتائج 1'].max_row,1);self.assertEqual(wb['النتائج 2'].max_row,2)
        wb.close();response.close()
        comparison=self.client.post('/api/compare',json={'db_id_a':'a','db_id_b':'b','key_col_a':'code','key_col_b':'code','filters_a':filters,'filters_b':filters})
        self.assertEqual(comparison.status_code,200,comparison.json)
        self.assertEqual(comparison.json['counts'],{'only_b':1})
        self.assertEqual(comparison.json['filtered_count_a'],0)
        response=self.client.post('/api/compare_export',json={'comparison_id':comparison.json['comparison_id']})
        wb=load_workbook(io.BytesIO(response.data));self.assertIn('أي خلية في الصف',[str(v).split(' — ')[0] for row in wb['ملخص المقارنة'].values for v in row if v]);wb.close();response.close()
        column=[{'col':'description','op':'not_contains','value':'المسجد'}]
        self.assertEqual(self.client.get('/api/search',query_string={'db_id':'a','filters':json.dumps(column)}).json['total'],1)
        self.assertEqual(self.client.get('/api/search',query_string={'db_id':'a','filters':json.dumps([{**filters[0],'value':''}])}).status_code,400)

    def test_saved_review_survives_cache_clear_and_source_change(self):
        result=self.client.post('/api/compare',json={'db_id_a':'a','db_id_b':'b','key_col_a':'description','key_col_b':'description','mode':'similarity','title':'مراجعة المضخات'}).json
        token=result['comparison_id'];rid=result['items'][0]['id']
        self.assertEqual(self.client.post('/api/compare_decision',json={'comparison_id':token,'id':rid,'decision':'accepted','note':'=مقبول بعد المراجعة'}).status_code,200)
        self.client.post('/api/compare_page',json={'comparison_id':token,'section':'accepted','offset':0})
        app.cleanup_comparisons();app.comparisons.clear()
        self.client.post('/api/update_cell',json={'db_id':'a','id':1,'col':'description','value':'مصدر معدل'})
        resumed=self.client.post('/api/compare_page',json={'comparison_id':token}).json
        self.assertEqual(resumed['section'],'accepted')
        self.assertEqual(resumed['items'][0]['note'],'=مقبول بعد المراجعة')
        self.assertEqual(resumed['items'][0]['row_a']['description'],'توريد مضخة مياه')
        self.assertEqual(self.client.get('/api/comparisons').json['sessions'][0]['title'],'مراجعة المضخات')
        self.client.post('/api/compare_rename',json={'comparison_id':token,'title':'مراجعة نهائية'})
        self.assertEqual(self.client.get('/api/comparisons').json['sessions'][0]['title'],'مراجعة نهائية')
        response=self.client.post('/api/compare_export',json={'comparison_id':token})
        wb=load_workbook(io.BytesIO(response.data));self.assertEqual(wb['ملاحظات المراجعة']['F6'].data_type,'s');wb.close();response.close()
        self.assertEqual(self.client.post('/api/compare_delete',json={'comparison_id':'../config'}).status_code,400)
        self.assertTrue(Path(app.CONFIG_PATH).exists())


if __name__=='__main__':
    unittest.main()
