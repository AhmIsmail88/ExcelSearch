import json
import sqlite3
import tempfile
import unittest
from pathlib import Path

from comparison import Comparison


class ComparisonTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.engines = []

    def tearDown(self):
        for engine in self.engines:
            engine.close()
        self.tmp.cleanup()

    def db(self, name, rows):
        path = self.root / name
        conn = sqlite3.connect(path)
        conn.execute('CREATE TABLE data(id INTEGER PRIMARY KEY,code TEXT,description TEXT,qty TEXT)')
        conn.executemany('INSERT INTO data(code,description,qty) VALUES(?,?,?)', rows)
        conn.commit()
        conn.close()
        return path

    def compare(self, a, b, **kwargs):
        engine = Comparison(a,b,kwargs.pop('key','code'),kwargs.pop('key_b','code'),**kwargs)
        self.engines.append(engine)
        return engine

    def test_pages_counts_and_snapshot_survive_source_change(self):
        a=self.db('a',[(str(i),'حفر','1') for i in range(51)]+[('', 'blank', '0')])
        b=self.db('b',[(str(i),'وصف مختلف','2') for i in range(51)])
        engine=self.compare(a,b,query='حفر')
        page=engine.page()
        self.assertEqual(page['total'],51)
        self.assertEqual(len(page['items']),20)
        self.assertEqual(page['blank_a'],1)
        self.assertEqual(len(engine.page(offset=40)['items']),11)
        conn=sqlite3.connect(b);conn.execute('DELETE FROM data');conn.commit();conn.close()
        self.assertEqual(engine.page()['total'],51)

    def test_duplicate_keys_and_literal_arabic_codes(self):
        a=self.db('a',[('أ','same','1'),('ا','same','1')])
        b=self.db('b',[('ا','same','1')])
        self.assertEqual(self.compare(a,b).page()['counts'],{'identical':1,'only_a':1})
        c=self.db('c',[('x','same','1'),('x','same','1')])
        with self.assertRaisesRegex(ValueError,'مكرر'):
            self.compare(c,b)

    def test_suggestions_are_unconfirmed_and_one_to_one(self):
        a=self.db('a',[('1','توريد وتركيب مضخة مياه','1'),('2','توريد وتركيب مضخة مياه','2')])
        b=self.db('b',[('X','توريد وتركيب مضخة مياه','3')])
        engine=self.compare(a,b,key='description',key_b='description',mode='similarity')
        result=engine.page()
        self.assertEqual(result['counts'],{'suggested':2})
        first,second=result['items']
        self.assertEqual(first['score'],100)
        engine.decide(first['id'],'accepted')
        with self.assertRaisesRegex(ValueError,'بالفعل'):
            engine.decide(second['id'],'accepted')
        engine.decide(first['id'],'rejected')
        engine.decide(second['id'],'accepted')
        self.assertEqual(engine.page('accepted')['total'],1)
        output=self.root/'export.xlsx';engine.export(output)
        from openpyxl import load_workbook
        wb=load_workbook(output,read_only=True)
        rows=list(wb['المقارنة'].values)
        self.assertIn('مقبول يدويًا',[r[0] for r in rows])
        self.assertIn('مرفوض',[r[0] for r in rows])
        wb.close()

    def test_filters_are_applied_before_matching_and_export_is_styled(self):
        a=self.db('a',[('1','توريد مضخة مياه','ميكانيكا'),('2','توريد مضخة مياه','كهرباء')])
        b=self.db('b',[('X','توريد مضخة مياه','ميكانيكا'),('Y','توريد مضخة مياه','كهرباء')])
        filters=[{'col':'qty','op':'equals','value':'ميكانيكا'}]
        engine=self.compare(a,b,key='description',key_b='description',mode='similarity',filters_a=filters,filters_b=filters,name_a='جدول أ',name_b='جدول ب')
        result=engine.page()
        self.assertEqual(result['total'],1)
        self.assertEqual((result['source_count_a'],result['filtered_count_a']),(2,1))
        self.assertEqual(result['items'][0]['row_b']['qty'],'ميكانيكا')
        output=self.root/'styled.xlsx';engine.export(output)
        from openpyxl import load_workbook
        wb=load_workbook(output)
        self.assertEqual(wb.sheetnames,['ملخص المقارنة','المقارنة'])
        sheet=wb['المقارنة']
        self.assertEqual(sheet.freeze_panes,'A6')
        self.assertTrue(sheet.sheet_view.rightToLeft)
        self.assertEqual(sheet['A5'].fill.fgColor.rgb[-6:],'19354A')
        self.assertEqual(sheet['B6'].number_format,'0.0%')
        self.assertEqual(sheet['B6'].value,1)
        self.assertTrue(sheet.auto_filter.ref.startswith('A5:H'))
        self.assertGreaterEqual(sheet.row_dimensions[6].height,34)
        wb.close()

    def test_key_mode_can_compare_only_selected_value_columns(self):
        a=self.db('a',[('1','same description','1')])
        b=self.db('b',[('1','same description','99')])
        engine=self.compare(a,b,field_a='description',field_b='description')
        self.assertEqual(engine.page()['counts'],{'identical':1})
        output=self.root/'selected.xlsx';engine.export(output)
        from openpyxl import load_workbook
        wb=load_workbook(output,read_only=True)
        fields=[row[5] for row in list(wb['المقارنة'].values)[5:]]
        self.assertEqual(fields,['code','description'])
        wb.close()
        with self.assertRaises(ValueError):
            self.compare(a,b,filters_a=[{'col':'missing','value':'x'}])

    def test_cancel_snapshot_leaves_source_unchanged(self):
        a=self.db('a',[('1','text','1')]);b=self.db('b',[('1','text','2')])
        def cancel(_):
            raise InterruptedError('cancel')
        with self.assertRaises(InterruptedError):
            self.compare(a,b,progress=cancel)
        conn=sqlite3.connect(a)
        self.assertEqual(conn.execute('SELECT count(*) FROM data').fetchone()[0],1)
        conn.close()

    def test_compound_keys_distinguish_duplicates_without_collisions(self):
        a=self.db('compound-a',[('x','القسم الأول','1'),('x','القسم الثاني','2'),('a | b','c','3'),('a','b | c','4'),('z','','1')])
        b=self.db('compound-b',[('x','القسم الأول','1'),('x','القسم الثاني','9'),('a | b','c','3'),('a','b | c','4')])
        engine=self.compare(a,b,key_columns_a=['code','description'],key_columns_b=['code','description'])
        self.assertEqual(engine.page()['counts'],{'changed':1,'identical':3})
        self.assertEqual(engine.page()['blank_a'],1)
        self.assertEqual(engine.page()['items'][0]['row_a']['description'],'القسم الثاني')
        with self.assertRaises(ValueError):self.compare(a,b,key_columns_a=['code','code'],key_columns_b=['code','description'])
        with self.assertRaises(ValueError):self.compare(a,b,key_columns_a=['code','description'],key_columns_b=['code'])

    def test_persistent_snapshot_opens_in_fresh_process_and_exports(self):
        import comparison_sessions as sessions
        import subprocess,sys
        a=self.db('pa',[('1','توريد مضخة','1')]);b=self.db('pb',[('X','توريد مضخة جديدة','2')])
        engine=self.compare(a,b,key='description',key_b='description',mode='similarity')
        token='a'*32;folder=self.root/'sessions';sessions.save(engine,folder,token,'جلسة اختبار')
        rid=engine.page()['items'][0]['id'];engine.decide(rid,'accepted','مراجعة محفوظة')
        engine.close();a.unlink();b.unlink()
        code="import sys,comparison_sessions as s; e,_=s.restore(sys.argv[1],sys.argv[2]); p=e.page('accepted'); assert p['total']==1 and p['items'][0]['note']; e.export(sys.argv[3]); e.close()"
        subprocess.run([sys.executable,'-c',code,str(folder),token,str(self.root/'resumed.xlsx')],check=True)
        self.assertTrue((self.root/'resumed.xlsx').exists())

    def _year_db(self,name,rows):
        path=self.root/name
        conn=sqlite3.connect(path)
        conn.execute('CREATE TABLE data(id INTEGER PRIMARY KEY,code TEXT,description TEXT,"السنة_الاولى" TEXT,"السنة_الثانية" TEXT,"السنة_الثالثة" TEXT)')
        conn.executemany('INSERT INTO data(code,description,"السنة_الاولى","السنة_الثانية","السنة_الثالثة") VALUES(?,?,?,?,?)',rows)
        conn.commit();conn.close()
        return path

    def test_quantity_numeric_equality_block_and_filter(self):
        a=self._year_db('qa',[('1','مضخة مياه','10','',''),('2','صمام هواء','391','5','')])
        b=self._year_db('qb',[('X','مضخة مياه','10.0','2',''),('Y','صمام هواء','390.99999999999994','5','')])
        engine=self.compare(a,b,key='description',key_b='description',mode='similarity')
        page=engine.page()
        self.assertTrue(page['has_qty'])
        by_code={item['row_a']['code']:item for item in page['items']}
        first=by_code['1'];q=first['qty']
        self.assertEqual([c['delta'] for c in q['cols']],[0.0,None,None])
        self.assertTrue(q['changed'])
        self.assertEqual((q['a_total'],q['b_total'],q['delta_total']),(10.0,12.0,2.0))
        self.assertEqual((q['a_filled'],q['b_filled']),(1,2))
        second=by_code['2'];q2=second['qty']
        self.assertFalse(q2['changed'])
        self.assertEqual(q2['delta_total'],0.0)
        self.assertEqual(page['qty_counts'],{'suggested':1})
        only=engine.page(qty_only=True)
        self.assertTrue(only['qty_only'])
        self.assertEqual(only['total'],1)
        self.assertEqual(only['items'][0]['row_a']['code'],'1')

    def test_multiyear_totals_and_export_quantity_sheet(self):
        a=self._year_db('ma',[('1','بنود المكافحة','36','92','')])
        b=self._year_db('mb',[('X','بنود المكافحة','','','27')])
        engine=self.compare(a,b,key='description',key_b='description',mode='similarity')
        page=engine.page()
        item=page['items'][0];q=item['qty']
        self.assertEqual((q['a_total'],q['b_total'],q['delta_total']),(128.0,27.0,-101.0))
        self.assertTrue(q['changed'])
        output=self.root/'qty.xlsx';engine.export(output)
        from openpyxl import load_workbook
        wb=load_workbook(output,read_only=True)
        self.assertIn('الكميات',wb.sheetnames)
        rows=list(wb['الكميات'].values)
        self.assertTrue(any(len(r)>4 and r[2]=='الإجمالي' and r[3]==128.0 and r[4]==27.0 for r in rows))
        wb.close()

    def test_selected_years_detect_redistribution_and_export_one_row(self):
        a=self._year_db('ya',[('1','=وصف مضخة','10','',''),('2','بند آخر','0','5','')])
        b=self._year_db('yb',[('1','وصف معدل','','10',''),('2','تغير الوصف فقط','0','5','')])
        pairs=[{'a':c,'b':c} for c in ['السنة_الاولى','السنة_الثانية','السنة_الثالثة']]
        engine=self.compare(a,b,quantity_mode='yearly',quantity_pairs=pairs,display_columns_a=['description','code'])
        page=engine.page()
        self.assertEqual(page['counts'],{'changed':1,'identical':1})
        qty=page['items'][0]['qty']
        self.assertEqual((qty['a_total'],qty['b_total'],qty['delta_total']),(10,10,0))
        self.assertIn('توزيع السنوات مختلف',qty['status'])
        self.assertIsNone(qty['cols'][0]['delta'])
        output=self.root/'wide.xlsx';engine.export(output)
        from openpyxl import load_workbook
        wb=load_workbook(output);sheet=wb['المقارنة']
        self.assertEqual(sheet.max_row,7)
        self.assertEqual(sheet['E6'].value,'=وصف مضخة')
        self.assertEqual(sheet['E6'].data_type,'s')
        self.assertEqual(sheet['H6'].value,None)
        self.assertEqual(sheet['G7'].value,0)
        self.assertEqual(sheet.freeze_panes,'G6')
        self.assertTrue(sheet.auto_filter.ref.startswith('A5:'))
        wb.close()
        total=self.compare(a,b,quantity_mode='total',quantity_pairs=pairs)
        self.assertEqual(total.page()['counts'],{'identical':2})
        self.assertIn('توزيع السنوات مختلف',total.page('identical')['items'][0]['qty']['status'])

    def test_custom_years_and_missing_invalid_quantities(self):
        a=self._year_db('ca',[('1','pump','10','',''),('2','blank','','',''),('3','zero','0','',''),('4','bad','NaN','','')])
        b=self._year_db('cb',[('1','pump','','10',''),('2','blank','','',''),('3','zero','','',''),('4','bad','','5','')])
        pairs=[{'a':'السنة_الاولى','b':'السنة_الثانية'}]
        engine=self.compare(a,b,quantity_mode='custom',quantity_pairs=pairs)
        self.assertEqual(engine.page()['counts'],{'changed':3,'identical':1})
        self.assertEqual(engine.page('identical')['items'][0]['qty']['delta_total'],0)
        by_key={item['key']:item['qty'] for item in engine.page()['items']}
        self.assertIsNone(by_key['2']['a_total'])
        self.assertEqual(by_key['3']['a_total'],0)
        self.assertIsNone(by_key['3']['delta_total'])
        self.assertIn('غير رقمية',by_key['4']['status'])
        for opts in [dict(quantity_pairs=[]),dict(quantity_pairs=pairs+pairs),dict(quantity_pairs=[{'a':'missing','b':'السنة_الثانية'}]),dict(quantity_pairs=pairs,display_columns_a=['missing'])]:
            with self.assertRaises(ValueError):self.compare(a,b,quantity_mode='custom',**opts)

    def test_quantities_with_similarity_keep_manual_review_and_scope(self):
        a=self._year_db('sqa',[('1','مضخة مياه','10','',''),('2','مضخة مياه','90','','')])
        b=self._year_db('sqb',[('X','مضخة مياه','','10','')])
        engine=self.compare(a,b,key='description',key_b='description',mode='similarity',quantity_mode='custom',quantity_pairs=[{'a':'السنة_الاولى','b':'السنة_الثانية'}],filters_a=[{'col':'code','op':'equals','value':'1'}])
        self.assertEqual(engine.page()['counts'],{'suggested':1})
        self.assertEqual(engine.page()['items'][0]['qty']['delta_total'],0)
        engine.decide(engine.page()['items'][0]['id'],'accepted')
        self.assertEqual(engine.page('accepted')['total'],1)

    def test_suggestion_totals_exclude_unreviewed_duplicate_candidates(self):
        a=self._year_db('sa',[('1','مضخة مياه','10','','')])
        b=self._year_db('sb',[('X','مضخة مياه','20','',''),('Y','مضخة مياه','30','','')])
        engine=self.compare(a,b,key='description',key_b='description',mode='similarity')
        from openpyxl import load_workbook
        for accepted in (False,True):
            if accepted: engine.decide(engine.page()['items'][0]['id'],'accepted')
            output=self.root/'totals.xlsx';engine.export(output)
            wb=load_workbook(output,read_only=True)
            row=list(wb['الكميات'].values)[3]
            self.assertEqual(row[1:3],(10,20) if accepted else (0,0))
            wb.close()


if __name__=='__main__':
    unittest.main()
