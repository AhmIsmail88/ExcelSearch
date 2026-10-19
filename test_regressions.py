import sqlite3
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
import excel_db_core as core


class RegressionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def workbook(self, name, rows):
        path = self.root / (name + '.xlsx')
        wb = Workbook()
        for row in rows:
            wb.active.append(row)
        wb.save(path)
        wb.close()
        return path

    def test_large_text_batches_search_and_sync(self):
        text = 'وصف طويل لأعمال المشروع ' * 600
        rows = [['code', 'description']] + [[str(i), text] for i in range(1100)]
        path = self.workbook('large', rows)
        db = self.root / 'large.sqlite'
        self.assertEqual(core.build_db(path, db, unique_col='code', header_row=1)['added'], 1100)
        found, total = core.search(db, 'المشروع', limit=2)
        self.assertEqual(total, 1100)
        self.assertEqual(found[0]['description'], text)
        self.assertEqual(core.search(db, 'المشروع', offset=1100)[1], 1100)
        core.update_cell(db, 1, 'description', 'local edit')
        core.clear_override(db, 1, 'description')
        self.assertFalse(core.sync_db(db)['skipped'])
        with sqlite3.connect(db) as conn:
            self.assertEqual(conn.execute('SELECT description FROM data WHERE id=1').fetchone()[0], text)
        conn.close()

    def test_failed_build_preserves_existing_database(self):
        path = self.workbook('duplicates', [['code'], ['a'], ['a']])
        db = self.root / 'existing.sqlite'
        db.write_bytes(b'original database')
        with self.assertRaises(ValueError):
            core.build_db(path, db, unique_col='code', header_row=1)
        self.assertEqual(db.read_bytes(), b'original database')
        self.assertEqual(list(self.root.glob('.index-*')), [])

    def test_data_beyond_last_header_is_not_truncated(self):
        path = self.workbook('wide', [['code'], ['a', 'extra value']])
        db = self.root / 'wide.sqlite'
        core.build_db(path, db, header_row=1)
        self.assertEqual(core.search(db, 'extra')[1], 1)
        self.assertEqual(core.sync_db(db, force=True)['unchanged'], 1)

    def test_sparse_far_columns_are_kept_without_empty_gaps(self):
        path = self.root / 'sparse.xlsx'
        wb = Workbook()
        ws = wb.active
        ws.append(['code', 'description'])
        ws.append(['a', 'normal text'])
        ws.cell(2, 3000, 'faraway')
        ws.cell(1, 4000, 'header only')
        ws.cell(2, 5000, '   ')
        ws.cell(3, 16384).number_format = '0.00'
        wb.save(path)
        wb.close()
        db = self.root / 'sparse.sqlite'
        core.build_db(path, db, header_row=1)
        self.assertEqual(core.get_columns(db), ['code', 'description', 'Unnamed__2999', 'header_only'])
        self.assertEqual(core.search(db, 'faraway')[1], 1)
        self.assertEqual(core.sync_db(db, force=True)['unchanged'], 1)

    def test_real_column_limit_has_clear_error_and_no_database(self):
        conn = sqlite3.connect(':memory:')
        limit = conn.getlimit(sqlite3.SQLITE_LIMIT_COLUMN)
        conn.close()
        path = self.workbook('too_wide', [[f'col{i}' for i in range(limit)], ['value'] * limit])
        db = self.root / 'too_wide.sqlite'
        with self.assertRaisesRegex(ValueError, str(limit)):
            core.build_db(path, db, header_row=1)
        self.assertFalse(db.exists())
        self.assertEqual(list(self.root.glob('.index-*')), [])

    def test_older_streaming_databases_keep_blank_columns_on_sync(self):
        path = self.workbook('old', [['code', None, 'value'], ['a', None, 'hello']])
        from unittest.mock import patch
        original = core._worksheet_records
        with patch.object(core, '_worksheet_records', side_effect=lambda ws, idx, **kwargs: original(ws, idx, False)):
            db = self.root / 'old.sqlite'
            core.build_db(path, db, header_row=1)
        conn = sqlite3.connect(db)
        conn.execute("DELETE FROM _config WHERE key='compact_columns'")
        conn.commit()
        conn.close()
        self.assertEqual(core.sync_db(db, force=True)['unchanged'], 1)

    def test_explicit_column_boundary_persists_in_sync(self):
        path = self.workbook('bounded', [['code', 'text'], ['a','retained'] + ['outside'] * 2500])
        db = self.root/'bounded.sqlite'
        core.build_db(path, db, header_row=1, last_column='B')
        self.assertEqual(core.get_columns(db), ['code','text'])
        self.assertEqual(core.search(db, 'retained')[1], 1)
        self.assertEqual(core.search(db, 'outside')[1], 0)
        self.assertEqual(core.sync_db(db, force=True)['unchanged'], 1)
        for value in ('XFE', '-1', '0', 'B:Z'):
            with self.assertRaises(ValueError):
                core.build_db(path, db, last_column=value)

    def test_overflow_error_suggests_last_real_column(self):
        conn = sqlite3.connect(':memory:')
        limit = conn.getlimit(sqlite3.SQLITE_LIMIT_COLUMN)
        conn.close()
        path = self.root / 'suggest.xlsx'
        wb = Workbook()
        ws = wb.active
        ws.append(['code', 'text', 'qty', 'price', 'notes'])
        for i in range(10):
            ws.append([f'c{i}', 'x', '1', '2', 'n'])
        for c in range(100, 100 + limit + 60):
            ws.cell(5, c, 'stray')
        wb.save(path)
        wb.close()
        db = self.root / 'suggest.sqlite'
        with self.assertRaisesRegex(ValueError, r'E \(5\)') as ctx:
            core.build_db(path, db, header_row=1)
        self.assertIn('لم يتم حذف أي بيانات', str(ctx.exception))
        self.assertFalse(db.exists())

    def test_overflow_without_fitting_suggestion_has_no_hint(self):
        conn = sqlite3.connect(':memory:')
        limit = conn.getlimit(sqlite3.SQLITE_LIMIT_COLUMN)
        conn.close()
        width = limit + 60
        path = self.root / 'nosuggest.xlsx'
        wb = Workbook()
        ws = wb.active
        ws.append(['c'] * width)
        ws.append(['v'] * width)
        wb.save(path)
        wb.close()
        db = self.root / 'nosuggest.sqlite'
        with self.assertRaises(ValueError) as ctx:
            core.build_db(path, db, header_row=1)
        self.assertIn('الحد المتاح', str(ctx.exception))
        self.assertNotIn('تلميح', str(ctx.exception))

    def test_reserved_duplicate_headers_and_blank_rows_sync(self):
        path = self.workbook('headers', [['id', 'same', 'same', 'A', 'a'], ['x','1','2','3','4'], [], ['y','5','6','7','8']])
        db = self.root / 'headers.sqlite'
        core.build_db(path, db, unique_col='id', header_row=1)
        self.assertEqual(core.get_columns(db), ['id_2','same','same_2','A','a_2'])
        self.assertEqual(core.sync_db(db, force=True)['unchanged'], 2)

    def test_comparison_rejects_duplicates_and_ignores_metadata(self):
        path = self.workbook('compare', [['code','value'], ['a','hello'], ['b','world']])
        a, b = self.root/'a.sqlite', self.root/'b.sqlite'
        for db in (a,b):
            core.build_db(path, db, header_row=1)
        core.update_cell(a, 1, 'value', 'hello')
        self.assertEqual(core.compare_databases(a,b,'code','code')['matched_identical'], 2)
        core.update_cell(a, 2, 'code', 'a')
        with self.assertRaises(ValueError):
            core.compare_databases(a,b,'code','code')

    def test_filters_combine_with_query_and_treat_symbols_literally(self):
        path = self.workbook('filters', [['code','description','category'],
                                        ['A','أعمال الحفر','مدني'],
                                        ['B','أعمال الحفر','كهرباء'],
                                        ['C','أعمال الطلاء','مدني'],
                                        ['D','100%_literal',None]])
        db = self.root/'filters.sqlite'
        core.build_db(path, db, header_row=1)
        filters = [{'col':'category','op':'equals','value':'مدني'}]
        rows,total = core.query_rows(db,'الحفر',filters)
        self.assertEqual((total,rows[0]['code']),(1,'A'))
        self.assertEqual(core.query_rows(db,'الحفر',filters,offset=1),( [],1))
        filters.append({'col':'description','op':'contains','value':'اعمال'})
        self.assertEqual(core.query_rows(db,'',filters)[1],2)
        self.assertEqual(core.query_rows(db,filters=[{'col':'description','value':'%_'}])[1],1)
        self.assertEqual(core.query_rows(db,filters=[{'col':'category','op':'empty'}])[1],1)
        for bad in ([{'col':'unknown'}],[{'col':'code','op':'SQL'}], {'col':'code'}):
            with self.assertRaises(ValueError):
                core.query_rows(db,filters=bad)

    def test_comparison_keyword_keeps_counterpart_with_changed_description(self):
        a,b=self.root/'qa.sqlite',self.root/'qb.sqlite'
        pa=self.workbook('qa',[['code','text'],['1','أعمال الحفر'],['2','دهان'],['3','حفر جديد']])
        pb=self.workbook('qb',[['code','text'],['1','وصف معدل'],['2','دهان']])
        core.build_db(pa,a,header_row=1)
        core.build_db(pb,b,header_row=1)
        result=core.compare_databases(a,b,'code','code',query='حفر')
        self.assertEqual(result['matched_changed'],1)
        self.assertEqual(result['only_in_a_count'],1)
        self.assertEqual(result['only_in_b_count'],0)
        self.assertEqual(result['matched_identical'],0)
        self.assertEqual(result['changed'][0]['row_b']['text'],'وصف معدل')

    def test_sync_blocks_positional_local_edits_after_reordering(self):
        path=self.workbook('reorder',[['code','text'],['1','first'],['2','second']])
        db=self.root/'reorder.sqlite'
        core.build_db(path,db,header_row=1)
        core.update_cell(db,1,'text','protected')
        self.workbook('reorder',[['code','text'],['2','second'],['1','first']])
        with self.assertRaisesRegex(ValueError,'مفتاح فريد'):
            core.sync_db(db)
        self.assertEqual(core.query_rows(db)[0][0]['text'],'protected')

    def test_streaming_sync_preserves_keyed_edits_and_rolls_back_cancel(self):
        path=self.workbook('safe',[['code','text'],['1','first'],['2','second']])
        db=self.root/'safe.sqlite';core.build_db(path,db,header_row=1,unique_col='code')
        core.update_cell(db,1,'text','protected')
        self.workbook('safe',[['code','text'],['2','changed'],['1','first'],['3','new']])
        def cancel(message):
            if message=='حفظ المزامنة...':
                raise InterruptedError('cancel')
        with self.assertRaises(InterruptedError):
            core.sync_db(db,progress=cancel)
        self.assertEqual(core.get_stats(db)['row_count'],2)
        result=core.sync_db(db)
        self.assertEqual(result['added'],1)
        rows,_=core.query_rows(db)
        self.assertEqual(rows[0]['text'],'protected')
        self.assertEqual(rows[1]['text'],'changed')


if __name__ == '__main__':
    unittest.main()
